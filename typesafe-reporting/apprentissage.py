"""Boucle d'apprentissage supervisée par les analystes (le modèle n'est pas réentraîné).

L'API ne permet pas de modifier les poids du modèle Jev. Ce sont ses consignes et nos seuils qui apprennent :

1. --modele-apprentissage : après un enrichissement, écrit sortie/a_valider.xlsx. On y trouve les dossiers
   commentés, ceux en revue analyste en premier, avec les jugements proposés (motif, orientation, statut,
   action) et leur confiance. La colonne « Réponse validée » propose une liste déroulante.
2. Les analystes valident ou corrigent ; une case vide n'est pas prise en compte.
3. --apprentissage a_valider.xlsx : les réponses validées rejoignent sortie/memoire_apprentissage.json.
   - Apprentissage par l'exemple : à chaque appel suivant, jusqu'à 2 exemples validés par option
     (12 au plus) accompagnent les instructions structurées de la question.
   - Recalibrage : dès 10 réponses validées, le seuil de confiance de la question devient le plus petit
     seuil (plancher 0,40) donnant au moins 90 % d'accord avec les analystes sur les cas au-dessus du seuil.
     En dessous, le dossier part en revue analyste.
   - La version de la mémoire entre dans la clé du cache : de nouveaux exemples déclenchent une relecture.

Les commentaires conservés sont masqués : le nom du client et celui du gestionnaire sont remplacés avant
stockage, car les exemples partent ensuite vers l'API.
"""
import hashlib
import json
import os
from datetime import datetime
from pathlib import Path

ICI = Path(__file__).resolve().parent
MEMOIRE = Path(os.environ.get("TYPESAFE_MEMOIRE") or ICI / "sortie" / "memoire_apprentissage.json")
QUESTIONS = ("motif", "routage", "statut_propose", "action_recommandee")
LIBELLES_Q = {"motif": "Motif de l'impayé", "routage": "Orientation du dossier", "statut_propose": "Statut de suivi",
              "action_recommandee": "Action recommandée"}
PAR_OPTION, MAXI, MIN_CALIBRAGE, CIBLE, PLANCHER = 2, 12, 10, 0.90, 0.40

_mem = None


def charger(force=False):
    global _mem
    if _mem is None or force:
        _mem = json.loads(MEMOIRE.read_text(encoding="utf-8")) if MEMOIRE.exists() else {"exemples": [], "seuils": {}}
    return _mem


def version():
    """Empreinte des exemples validés : vide tant que la mémoire est vide (clés de cache inchangées)."""
    ex = charger()["exemples"]
    if not ex:
        return ""
    return hashlib.sha1(json.dumps([[e["question"], e["commentaire"], e["valide"]] for e in ex],
                                   ensure_ascii=False).encode()).hexdigest()[:10]


def exemples(q):
    """Exemples validés pour la question : diversifiés par option, les plus récents d'abord."""
    vus, out = {}, []
    for e in sorted((e for e in charger()["exemples"] if e["question"] == q), key=lambda e: e["date"], reverse=True):
        if vus.get(e["valide"], 0) < PAR_OPTION and len(out) < MAXI:
            vus[e["valide"]] = vus.get(e["valide"], 0) + 1
            out.append({"commentaire": e["commentaire"], "reponse_attendue": e["valide"]})
    return out


def enrichir(questions):
    """Joint les exemples validés aux instructions (y compris aux variantes à options inversées)."""
    if not charger()["exemples"]:
        return questions
    out = dict(questions)
    for k, q in questions.items():
        base = k[:-4] if k.endswith("_inv") else k
        if base not in QUESTIONS or q.get("type") != "choice":
            continue
        ex = exemples(base)
        if ex:
            out[k] = dict(q, instructions={
                "question": q["instructions"],
                "exemples_valides_par_les_analystes": ex,
                "consigne": "Ces lectures ont été validées par des analystes risque de la banque : appliquez la même "
                            "lecture aux situations comparables, sans les recopier quand la situation diffère."})
    return out


def seuil(q, defaut):
    s = charger()["seuils"].get(q)
    return s["seuil"] if s else defaut


def calibrer(mem):
    seuils = {}
    for q in QUESTIONS:
        ex = [e for e in mem["exemples"] if e["question"] == q and e.get("confiance") is not None]
        if len(ex) < MIN_CALIBRAGE:
            continue
        choix = None
        for t in [x / 100 for x in range(int(PLANCHER * 100), 96, 5)]:
            auto = [e for e in ex if e["confiance"] >= t]
            if len(auto) >= 5 and sum(e["propose"] == e["valide"] for e in auto) / len(auto) >= CIBLE:
                choix = (t, auto)
                break
        if choix:
            t, auto = choix
            seuils[q] = {"seuil": round(t, 2), "n": len(ex), "accord": round(sum(e["propose"] == e["valide"] for e in auto) / len(auto), 3),
                         "couverture": round(len(auto) / len(ex), 3)}
        else:
            seuils[q] = {"seuil": 0.95, "n": len(ex), "accord": None, "couverture": None,
                         "note": "accord insuffisant à tout seuil : revue analyste quasi systématique"}
    return seuils


def resume():
    m = charger()
    if not m["exemples"]:
        return "Mémoire vide : aucun exemple validé, seuils par défaut."
    par_q = {q: sum(e["question"] == q for e in m["exemples"]) for q in QUESTIONS}
    s = "; ".join(f"{LIBELLES_Q[q]} {v['seuil']:.2f}".replace(".", ",") + (f" (accord {v['accord']:.0%}, {v['n']} cas)" if v.get("accord") else f" ({v['n']} cas)")
                  for q, v in m["seuils"].items())
    return (f"{len(m['exemples'])} réponses validées (" + ", ".join(f"{LIBELLES_Q[q].lower()} {n}" for q, n in par_q.items() if n)
            + ")" + (f" ; seuils recalibrés : {s}" if s else " ; moins de 10 cas par question : seuils par défaut"))


# ---------------- fichier à valider et intégration des validations ----------------

def ecrire_a_valider(dossiers, resultats, cache, cle, libelles, chemin, masquer):
    """libelles : {question: {clé: libellé}} ; masquer(texte, d) retire nom du client et gestionnaire."""
    import xlsxwriter
    wb = xlsxwriter.Workbook(str(chemin))
    ws, li = wb.add_worksheet("À valider"), wb.add_worksheet("_listes")
    hd = wb.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#001B4D", "bottom": 2, "bottom_color": "#C8A951",
                        "text_wrap": True, "valign": "vcenter", "font_name": "Calibri", "font_size": 10})
    tt = wb.add_format({"bold": True, "font_color": "#001B4D", "font_size": 15, "font_name": "Calibri"})
    st = wb.add_format({"italic": True, "font_color": "#64748B", "font_size": 9, "font_name": "Calibri", "text_wrap": True})
    ce = wb.add_format({"text_wrap": True, "valign": "top", "font_size": 9, "font_name": "Calibri", "border": 1, "border_color": "#DBE6F7"})
    pc = wb.add_format({"num_format": "0%", "valign": "top", "align": "center", "font_size": 9, "border": 1, "border_color": "#DBE6F7"})
    va = wb.add_format({"bg_color": "#FFF7E0", "valign": "top", "font_size": 9, "border": 1, "border_color": "#C8A951", "bold": True})
    rv = wb.add_format({"text_wrap": True, "valign": "top", "font_size": 9, "font_color": "#DC2626", "border": 1, "border_color": "#DBE6F7"})
    ws.write(0, 0, "APPRENTISSAGE SUPERVISÉ — RÉPONSES À VALIDER PAR LES ANALYSTES", tt)
    ws.merge_range(1, 0, 1, 9, "Validez ou corrigez la colonne « Réponse validée » (liste déroulante), laissez vide en cas de doute. "
                   "Puis : python3 impayes_typesafe.py --apprentissage a_valider.xlsx. Les réponses validées deviennent des exemples "
                   "joints aux consignes du moteur et servent à recalibrer les seuils de confiance. Le modèle n'est pas réentraîné.", st)
    ws.set_row(1, 42)
    cols = ["Référence", "Client", "Commentaire du gestionnaire", "Question", "Réponse proposée", "Confiance",
            "Réponse validée", "Remarque analyste", "Revue analyste (motifs)", "_question", "_propose"]
    for c, h in enumerate(cols):
        ws.write(3, c, h, hd)
    for j, q in enumerate(QUESTIONS):
        li.write(0, j, q)
        for i, lib in enumerate(libelles[q].values()):
            li.write(1 + i, j, lib)
    li.hide()
    ordre = sorted(zip(dossiers, resultats), key=lambda x: (not x[1]["revue"], x[0]["ref"]))
    y = 4
    for d, r in ordre:
        if not d["commentaire"] or cle(d) not in cache:
            continue
        A = cache[cle(d)]["answers"]
        for j, q in enumerate(QUESTIONS):
            if q not in A:
                continue
            ch, cf = A[q]["choice"], A[q]["confidence"]
            ws.write(y, 0, d["ref"], ce); ws.write(y, 1, d["client"], ce); ws.write(y, 2, masquer(d["commentaire"], d), ce)
            ws.write(y, 3, LIBELLES_Q[q], ce); ws.write(y, 4, libelles[q].get(ch, ch), ce); ws.write(y, 5, cf, pc)
            ws.write_blank(y, 6, None, va); ws.write_blank(y, 7, None, ce); ws.write(y, 8, " ; ".join(r["revue"]), rv)
            ws.write(y, 9, q); ws.write(y, 10, ch)
            n = len(libelles[q])
            col = "ABCD"[j]
            ws.data_validation(y, 6, y, 6, {"validate": "list", "source": f"=_listes!${col}$2:${col}${n + 1}"})
            ws.set_row(y, 48)
            y += 1
    ws.set_column(0, 0, 13); ws.set_column(1, 1, 22); ws.set_column(2, 2, 60); ws.set_column(3, 3, 18)
    ws.set_column(4, 4, 26); ws.set_column(5, 5, 9); ws.set_column(6, 6, 28); ws.set_column(7, 7, 26); ws.set_column(8, 8, 34)
    ws.set_column(9, 10, None, None, {"hidden": True})
    ws.freeze_panes(4, 2); ws.autofilter(3, 0, max(4, y - 1), 8)
    ws.set_landscape(); ws.set_paper(9); ws.fit_to_pages(1, 0); ws.repeat_rows(3); ws.set_margins(0.4, 0.4, 0.5, 0.5)
    wb.close()
    return y - 4


def integrer(fichier, libelles, masquer_ligne):
    """Lit le fichier validé ; masquer_ligne(commentaire, client) retire les noms avant stockage."""
    import openpyxl
    inv = {q: {v: k for k, v in libelles[q].items()} for q in QUESTIONS}
    ws = openpyxl.load_workbook(fichier, read_only=True, data_only=True)["À valider"]
    mem = charger(force=True)
    connus = {(e["question"], e["commentaire"]): e for e in mem["exemples"]}
    ajout = maj = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row or not row[9] or not row[6]:
            continue
        q, valide = row[9], inv.get(row[9], {}).get(str(row[6]).strip())
        if q not in QUESTIONS or not valide:
            continue
        com = masquer_ligne(str(row[2] or ""), str(row[1] or ""))
        e = {"question": q, "commentaire": com, "propose": row[10], "confiance": float(row[5]) if row[5] is not None else None,
             "valide": valide, "remarque": row[7] or "", "date": datetime.now().isoformat(timespec="seconds")}
        if (q, com) in connus:
            connus[(q, com)].update(e); maj += 1
        else:
            mem["exemples"].append(e); connus[(q, com)] = e; ajout += 1
    mem["seuils"] = calibrer(mem)
    mem["maj"] = datetime.now().isoformat(timespec="seconds")
    MEMOIRE.parent.mkdir(exist_ok=True)
    MEMOIRE.write_text(json.dumps(mem, ensure_ascii=False, indent=1), encoding="utf-8")
    charger(force=True)
    return ajout, maj
