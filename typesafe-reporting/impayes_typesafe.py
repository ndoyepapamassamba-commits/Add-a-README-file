#!/usr/bin/env python3
"""Impayés 30-90j / Plan d'actions : couche TypeSafe au retour du fichier gestionnaires.

Chaîne de traitement
  1. Lecture du fichier retourné par les gestionnaires (feuille « Plan d'actions »),
     ou d'une copie anonymisée de démonstration (--demo).
  2. Code (règles)      : classe et date de bascule en douteux (> 90 jours), date promise
                          extraite du commentaire, écart avec la date de bascule.
  3. TypeSafe (Jev)     : une requête par dossier ; jugements uniquement, jamais de chiffre.
                          Commentaire : motif, promesse datée, crédibilité, incohérence,
                          statut du suivi proposé. Nom du client : public / parapublic,
                          intra-groupe Ecobank, segment douteux.
  4. Politique (code)   : statut retenu, renvoi en revue analyste si la confiance est faible.
  5. Restitution        : classeur Excel BLUE ECOBANK avec les feuilles Retours gestionnaires,
                          Promesses vs bascules, Contrôle qualité, Méthodologie TypeSafe et
                          _TYPESAFE (réponses brutes, masquée) que l'app HTML offline relit.

Usage
  python3 impayes_typesafe.py --demo                    # copie anonymisée + commentaires fictifs
  python3 impayes_typesafe.py --demo --depuis-cache     # réapplique les seuils sans appel API
  python3 impayes_typesafe.py --fichier retour.xlsx     # fichier réel (après accord Conformité)
  python3 impayes_typesafe.py --lot Lot_TypeSafe_Impayes.json [--anonymiser]
                                                        # lot exporté par APEX (salle JEV) ; le classeur
                                                        # produit se recharge dans APEX (« Retours TypeSafe »)

La clé API n'est jamais écrite dans un fichier : TYPESAFE_API_KEY est lue dans l'environnement
du poste connecté (ou fournie par le relais serveur).
"""

import argparse
import calendar
import csv
import json
import os
import random
import re
import time
import unicodedata
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from pathlib import Path

import xlsxwriter

import apprentissage as app
import typesafe_avance as ava

API_URL = "https://api.typesafe.ai/v1/systemone"
MODELE = "jev-latest"
ICI = Path(__file__).resolve().parent
SORTIE = ICI / "sortie"
CACHE = Path(os.environ.get("TYPESAFE_CACHE") or SORTIE / "impayes_reponses_brutes.json")
DEMO_CSV = ICI / "impayes_demo_anonyme.csv"

# --------------------------------------------------------------------------- #
# 1. Règles déterministes (code). Seuils simplifiés pour la démonstration.
# --------------------------------------------------------------------------- #

SEUIL_DOUTEUX = 90          # jours d'impayés au-delà desquels le dossier bascule en douteux


def classe_impaye(jours):
    if jours > SEUIL_DOUTEUX:
        return "Douteux"
    if jours > 60:
        return "Sensible 61-90j"
    return "Sensible 31-60j"


def date_bascule(arrete, jours):
    """Date à laquelle le dossier atteint 91 jours d'impayés si rien n'est réglé."""
    return arrete + timedelta(days=max(0, SEUIL_DOUTEUX + 1 - jours))


MOIS = {m: i for i, m in enumerate(
    "janvier fevrier mars avril mai juin juillet aout septembre octobre novembre decembre".split(), 1)}


def sans_accent(t):
    return "".join(c for c in unicodedata.normalize("NFD", t or "") if unicodedata.category(c) != "Mn").lower()


def fin_de_mois(a, m):
    return date(a, m, calendar.monthrange(a, m)[1])


def date_promise(texte, arrete):
    """Extraction par le code de la date de régularisation citée dans le commentaire.

    TypeSafe dit seulement si une date est promise ; la date elle-même est lue ici,
    pour rester traçable et reproductible.
    """
    t = sans_accent(texte)
    if not t:
        return None
    m = re.search(r"\b(\d{1,2})[/.-](\d{1,2})(?:[/.-](\d{2,4}))?\b", t)
    if m:
        j, mo = int(m.group(1)), int(m.group(2))
        a = int(m.group(3) or arrete.year)
        a += 2000 if a < 100 else 0
        try:
            d = date(a, mo, j)
            return d if m.group(3) or d >= arrete else date(a + 1, mo, j)
        except ValueError:
            pass
    m = re.search(r"\b(\d{1,2}) (" + "|".join(MOIS) + r")(?: (\d{4}))?", t)
    if m:
        mo = MOIS[m.group(2)]
        a = int(m.group(3) or (arrete.year if mo >= arrete.month else arrete.year + 1))
        try:
            return date(a, mo, int(m.group(1)))
        except ValueError:
            pass
    m = re.search(r"fin (?:du mois (?:d'|de )?)?(" + "|".join(MOIS) + r")", t)
    if m:
        mo = MOIS[m.group(1)]
        return fin_de_mois(arrete.year if mo >= arrete.month else arrete.year + 1, mo)
    if re.search(r"fin (du|de) mois", t):
        return fin_de_mois(arrete.year, arrete.month)
    if "fin du mois prochain" in t or "fin de mois prochain" in t:
        mo = arrete.month % 12 + 1
        return fin_de_mois(arrete.year + (mo == 1), mo)
    m = re.search(r"(?:sous|dans|d'ici) (\d{1,3}) jours", t)
    if m:
        return arrete + timedelta(days=int(m.group(1)))
    if "semaine prochaine" in t:
        return arrete + timedelta(days=7 + 4 - arrete.weekday())
    if "cette semaine" in t:
        return arrete + timedelta(days=max(0, 4 - arrete.weekday()))
    m = re.search(r"\b(?:le|au|avant le) (\d{1,2})\b", t)
    if m:
        j = int(m.group(1))
        a, mo = arrete.year, arrete.month
        if j <= arrete.day:
            mo = mo % 12 + 1
            a += mo == 1
        try:
            return date(a, mo, j)
        except ValueError:
            pass
    return None


# --------------------------------------------------------------------------- #
# 2. Questions TypeSafe : jugements étroits, jamais de montant ni de ratio.
# --------------------------------------------------------------------------- #

Q_COMMENTAIRE = {
    "motif": {
        "type": "choice",
        "instructions": "Quelle est la cause principale de l'impayé décrite dans `dossier.commentaire_gestionnaire` ?",
        "criteria": {
            "retard_technique": "Virement en cours, domiciliation, prélèvement ou délai administratif ; la capacité de paiement n'est pas en cause",
            "tension_tresorerie": "Décalage de trésorerie temporaire, client payeur ou État en retard, activité maintenue",
            "perte_revenus": "Perte durable de revenus : contrat perdu, activité en baisse, perte d'emploi, salaire non domicilié",
            "litige_juridique": "Litige commercial, fiscal ou judiciaire, saisie ou blocage de comptes",
            "sante_deces": "Maladie, hospitalisation ou décès de l'emprunteur ou d'un dirigeant",
            "fraude_detournement": "Suspicion de fraude, de détournement ou de faux documents",
            "client_injoignable": "Client injoignable, absent ou parti sans information",
            "regularisation_annoncee": "Le gestionnaire annonce que l'impayé est déjà réglé",
            "aucun_motif": "Le commentaire ne donne aucune cause exploitable",
        },
    },
    "promesse_datee": {
        "type": "noul",
        "instructions": "Le commentaire cite-t-il une date ou un délai précis de régularisation (le 15, fin octobre, sous 10 jours, cette semaine…) ?",
        "criteria": {
            "true": "Une date ou un délai précis de régularisation est cité",
            "false": "Aucune date : promesse vague ou absente",
        },
    },
    "credibilite": {
        "type": "score",
        "instructions": "Quelle est la crédibilité de la régularisation annoncée dans `dossier.commentaire_gestionnaire`, au vu de la cause et des preuves évoquées ?",
        "criteria": [
            "Nulle : aucune promesse, ou promesse contredite par la situation décrite",
            "Faible : promesse vague, sans source de fonds identifiée",
            "Moyenne : source de fonds identifiée mais incertaine",
            "Bonne : source de fonds identifiée et engagement daté",
            "Très bonne : preuve tangible (ordre de virement, chèque, domiciliation rétablie, mandat de l'État)",
        ],
    },
    "incoherence": {
        "type": "noul",
        "instructions": {
            "question": "Le commentaire contredit-il les données du dossier ? Par exemple « régularisé », « à jour » ou « aucun impayé » alors que `dossier.jours_impayes` est supérieur à 30.",
            "rappel": "`dossier.jours_impayes` est observé dans le système à la date d'arrêté ; c'est la source de vérité.",
        },
        "criteria": {
            "true": "Le texte affirme une situation incompatible avec les données",
            "false": "Le texte est compatible avec les données",
        },
    },
    "statut_propose": {
        "type": "choice",
        "instructions": "Quel statut de suivi correspond le mieux à ce que décrit `dossier.commentaire_gestionnaire` ?",
        "criteria": {
            "a_faire": "Aucune action engagée ou commentaire sans contenu",
            "en_cours": "Action engagée (relance, visite, échange avec le client) sans engagement ferme",
            "promesse_reglement": "Le client s'est engagé à régler à une échéance",
            "regularise_a_verifier": "Le gestionnaire affirme que le règlement est fait : à vérifier dans le système",
            "restructuration": "Rééchelonnement ou restructuration demandé ou en cours",
            "escalade_contentieux": "Situation à transmettre au recouvrement, au juridique ou au contentieux",
        },
    },
    # --- capacités avancées : re-classement, routage, auto-cohérence, cohérence du plan ---
    "urgence": {
        "type": "score",
        "instructions": "Quelle est l'urgence d'une intervention de la banque sur ce dossier, au vu de `dossier.commentaire_gestionnaire` et de `dossier.jours_impayes` ?",
        "criteria": [
            "Aucune : incident réglé ou purement technique",
            "Faible : suivi normal suffisant",
            "Modérée : relance structurée nécessaire",
            "Élevée : intervention sous 7 jours",
            "Critique : intervention immédiate, risque de perte",
        ],
    },
    "routage": {
        "type": "choice",
        "instructions": "Vers quel traitement ce dossier doit-il être orienté, au vu de `dossier.commentaire_gestionnaire` ?",
        "criteria": {
            "gestionnaire": "Le gestionnaire peut régulariser seul (relance, ordre de virement, domiciliation)",
            "recouvrement_amiable": "Relance structurée par le recouvrement amiable",
            "restructuration_credit": "Étude de rééchelonnement ou de restructuration par le crédit",
            "juridique_contentieux": "Litige, fraude, saisie ou client disparu : juridique ou contentieux",
            "analyste_risque": "Situation ambiguë ou contradictoire : revue par un analyste risque",
        },
    },
    "motif_durable": {
        "type": "noul",
        "instructions": "La difficulté décrite dans `dossier.commentaire_gestionnaire` est-elle durable (perte de revenus, fraude, litige, client disparu) plutôt que ponctuelle ?",
        "criteria": {"true": "Difficulté durable ou structurelle", "false": "Difficulté ponctuelle, technique ou déjà réglée"},
    },
    "coherence_plan": {
        "type": "noul",
        "instructions": "Le plan d'action prévu `dossier.plan_action` traite-t-il la cause décrite dans `dossier.commentaire_gestionnaire` ?",
        "criteria": {"true": "Le plan répond à la cause décrite", "false": "Le plan ne traite pas la cause, ou aucun plan"},
    },
}
MOTIFS_DURABLES = {"perte_revenus", "fraude_detournement", "litige_juridique", "client_injoignable"}
LIB_ROUTAGE = {"gestionnaire": "Gestionnaire", "recouvrement_amiable": "Recouvrement amiable",
               "restructuration_credit": "Restructuration (crédit)", "juridique_contentieux": "Juridique / contentieux",
               "analyste_risque": "Analyste risque"}

Q_NOM = {
    "nature_entite": {
        "type": "choice",
        "instructions": "D'après `dossier.client`, de quel type d'entité s'agit-il ?",
        "criteria": {
            "publique": "Administration, ministère, collectivité locale ou établissement public de l'État",
            "parapublique": "Société nationale, société d'économie mixte, agence ou office public à caractère commercial",
            "privee": "Entreprise ou association privée",
            "particulier": "Personne physique",
        },
    },
    "intra_groupe": {
        "type": "noul",
        "instructions": "`dossier.client` désigne-t-il une entité du groupe Ecobank (filiale, ETI, fondation ou société liée Ecobank) ?",
    },
    "segment_douteux": {
        "type": "noul",
        "instructions": {
            "question": "Le segment déclaré `dossier.segment` paraît-il incompatible avec la nature du client `dossier.client` ?",
            "exemples": "Une personne physique classée Corporate, une grande société classée Particuliers, une administration classée PME.",
        },
        "criteria": {
            "true": "Le segment déclaré semble erroné au vu du nom",
            "false": "Le segment déclaré est plausible",
        },
    },
}


def cle(d):
    """Clé du cache : un commentaire modifié déclenche une nouvelle lecture TypeSafe."""
    import hashlib
    v = app.version()  # exemples validés par les analystes : une mémoire modifiée déclenche une relecture
    contexte = [d.get(k) or "" for k in ("secteur", "garantie", "taille", "trajectoire")]  # état transmis au modèle
    h = hashlib.sha1(json.dumps(["v3", d["client"], d["segment"], d["jours"], d["commentaire"], d.get("plan", "")]
                                + (contexte if any(contexte) else []) + ([v] if v else []),
                                ensure_ascii=False).encode()).hexdigest()
    return f"{d['ref']}|{h[:12]}"


def evaluer(d, essais=5):
    questions = dict(Q_NOM)
    if d["commentaire"]:
        questions.update(Q_COMMENTAIRE)
        questions.update(ava.questions_avancees(Q_COMMENTAIRE))
        questions.update(ava.questions_scenarios())
    etat = {"dossier": {
        "client": d["client"],
        "segment": d["segment"],
        "secteur": d.get("secteur") or "non renseigné",
        "garantie": d.get("garantie") or "non renseignée",
        "taille_exposition": d.get("taille") or "non renseignée",
        "trajectoire_impayes": d.get("trajectoire") or "non renseignée",
        "jours_impayes": d["jours"],
        "commentaire_gestionnaire": d["commentaire"],
        "plan_action": d.get("plan") or "aucun plan renseigné",
    }}
    if d.get("_sel"):
        etat = ava.etat_sel(etat)
        questions = {k: questions[k] for k in ("motif", "routage") if k in questions}
    questions = app.enrichir(questions)  # apprentissage par l'exemple (réponses validées par les analystes)
    corps = json.dumps({"model": MODELE, "state": etat, "questions": questions}).encode()
    entetes = {"Content-Type": "application/json"}
    if os.environ.get("TYPESAFE_API_KEY"):
        entetes["Authorization"] = f"Bearer {os.environ['TYPESAFE_API_KEY']}"
    for essai in range(essais):
        try:
            req = urllib.request.Request(API_URL, data=corps, headers=entetes)
            with urllib.request.urlopen(req, timeout=90) as rep:
                js = json.load(rep)
                return {"modele": MODELE, "version": js.get("model"), "usage": js.get("usage"),
                        "horodatage": datetime.now().isoformat(timespec="seconds"), "answers": js["answers"]}
        except urllib.error.HTTPError as e:
            if e.code in (429, 500, 502, 503, 529) and essai < essais - 1:
                time.sleep(2 ** essai)
                continue
            raise RuntimeError(f"{d['ref']} : HTTP {e.code} {e.read()[:300]!r}")
        except urllib.error.URLError:
            if essai < essais - 1:
                time.sleep(2 ** essai)
                continue
            raise


# --------------------------------------------------------------------------- #
# 3. Politique (code) : seuils explicites, réappliqués sans nouvel appel.
# --------------------------------------------------------------------------- #

SEUIL_SIGNAL = 0.6
SEUIL_CONFIANCE = 0.55
LIB_STATUT = {
    "a_faire": "À faire", "en_cours": "En cours", "promesse_reglement": "Promesse de règlement",
    "regularise_a_verifier": "Régularisé (à vérifier)", "restructuration": "Restructuration",
    "escalade_contentieux": "Escalade contentieux",
}
LIB_MOTIF = {
    "retard_technique": "Retard technique", "tension_tresorerie": "Tension de trésorerie",
    "perte_revenus": "Perte de revenus", "litige_juridique": "Litige juridique",
    "sante_deces": "Santé / décès", "fraude_detournement": "Fraude / détournement",
    "client_injoignable": "Client injoignable", "regularisation_annoncee": "Régularisation annoncée",
    "aucun_motif": "Aucun motif", None: "Sans retour",
}
LIB_NATURE = {"publique": "Publique", "parapublique": "Parapublique", "privee": "Privée", "particulier": "Particulier"}
LIB_CRED = ["Nulle", "Faible", "Moyenne", "Bonne", "Très bonne"]
# libellés proposés aux analystes dans le fichier à valider (listes déroulantes)
LIBELLES_APPRENTISSAGE = {"motif": {k: v for k, v in LIB_MOTIF.items() if k}, "routage": LIB_ROUTAGE,
                          "statut_propose": LIB_STATUT, "action_recommandee": ava.LIB_ACTION}


def analyser(d, brut, arrete):
    r = brut["answers"]
    out = dict(d)
    out["nature"] = LIB_NATURE[r["nature_entite"]["choice"]]
    out["conf_nature"] = r["nature_entite"]["confidence"]
    out["intra_groupe"] = r["intra_groupe"]["noul"] >= SEUIL_SIGNAL
    out["segment_douteux"] = r["segment_douteux"]["noul"] >= SEUIL_SIGNAL
    out["p_intra"] = r["intra_groupe"]["noul"]
    out["p_segment"] = r["segment_douteux"]["noul"]
    revue, qc = [], []
    if out["conf_nature"] < SEUIL_CONFIANCE:
        revue.append(f"nature de l'entité incertaine ({out['conf_nature']:.2f})")

    if not d["commentaire"]:
        out.update(motif=LIB_MOTIF[None], conf_motif=None, p_promesse=None, credibilite=None,
                   conf_cred=None, p_incoherence=None, statut_ts=None, conf_statut=None,
                   date_promise=None, ecart=None, position="Sans retour", statut_retenu="À faire")
        qc.append("Commentaire vide : relancer le gestionnaire")
    else:
        out["motif"] = LIB_MOTIF[r["motif"]["choice"]]
        out["conf_motif"] = r["motif"]["confidence"]
        out["p_promesse"] = r["promesse_datee"]["noul"]
        out["credibilite"] = int(round(r["credibilite"]["score"]))
        out["conf_cred"] = r["credibilite"]["confidence"]
        out["p_incoherence"] = r["incoherence"]["noul"]
        out["statut_ts"] = LIB_STATUT[r["statut_propose"]["choice"]]
        out["conf_statut"] = r["statut_propose"]["confidence"]

        dp = date_promise(d["commentaire"], arrete)
        out["date_promise"] = dp
        out["ecart"] = (dp - d["bascule"]).days if dp else None
        if dp is None:
            out["position"] = "Sans date"
        elif dp > d["bascule"]:
            out["position"] = "Après la bascule"
        else:
            out["position"] = "Avant la bascule"

        if out["conf_motif"] < app.seuil("motif", SEUIL_CONFIANCE):
            revue.append(f"motif incertain ({out['conf_motif']:.2f})")
        if out["conf_statut"] < app.seuil("statut_propose", SEUIL_CONFIANCE):
            revue.append(f"statut incertain ({out['conf_statut']:.2f})")
        if out["conf_cred"] < SEUIL_CONFIANCE:
            revue.append(f"crédibilité incertaine ({out['conf_cred']:.2f})")
        if out["p_incoherence"] >= SEUIL_SIGNAL:
            revue.append("commentaire incohérent avec les jours d'impayés")
        if (out["p_promesse"] >= SEUIL_SIGNAL) != (dp is not None):
            qc.append("Promesse datée selon TypeSafe mais date illisible par le code"
                      if dp is None else "Date lue par le code mais promesse jugée vague")
        if out["position"] == "Après la bascule":
            qc.append(f"Date promise {out['ecart']} j après la bascule en douteux")
        if dp and dp < arrete:
            qc.append("Date promise déjà dépassée à l'arrêté")
        if len(d["commentaire"]) < 25:
            qc.append("Commentaire trop court pour être exploitable")
        # re-classement, routage et contrôles d'auto-cohérence (TypeSafe pose la même question sous deux formes)
        if "urgence" in r:
            out["urgence"] = int(round(r["urgence"]["score"]))
            out["conf_urgence"] = r["urgence"]["confidence"]
            out["routage"] = LIB_ROUTAGE[r["routage"]["choice"]]
            out["conf_routage"] = r["routage"]["confidence"]
            out["p_durable"] = r["motif_durable"]["noul"]
            out["p_coherence_plan"] = r["coherence_plan"]["noul"]
            durable_choice = r["motif"]["choice"] in MOTIFS_DURABLES
            if (out["p_durable"] >= SEUIL_SIGNAL) != durable_choice and abs(out["p_durable"] - 0.5) > 0.15:
                revue.append("auto-cohérence : motif et durabilité divergent")
            if out["conf_routage"] < app.seuil("routage", SEUIL_CONFIANCE):
                revue.append(f"orientation incertaine ({out['conf_routage']:.2f})")
            if d.get("plan") and out["p_coherence_plan"] < 1 - SEUIL_SIGNAL:
                qc.append("Plan d'action APEX sans rapport avec la cause décrite")
        # capacités avancées : dates par composants, famille avec repli, action, sensibilité, contrôle d'ordre, scénarios
        if "famille_motif" in r:
            fam = r["famille_motif"]
            out["famille"] = ava.LIB_FAMILLE[fam["choice"]]
            if out["conf_motif"] < SEUIL_CONFIANCE and fam["confidence"] >= 0.9 and fam["choice"] != "indetermine":
                out["motif"] = out["famille"] + " (famille)"
                revue[:] = [x for x in revue if not x.startswith("motif incertain")]
            act = r["action_recommandee"]
            out["action"] = ava.LIB_ACTION[act["choice"]]
            out["conf_action"] = act["confidence"]
            out["p_sensible"] = r["donnees_sensibles"]["noul"]
            if out["p_sensible"] >= SEUIL_SIGNAL:
                qc.append("Donnée personnelle sensible dans le commentaire : à anonymiser avant diffusion")
            if r["motif_inv"]["choice"] != r["motif"]["choice"]:
                revue.append("contrôle d'ordre : le motif change quand les options sont inversées")
            if r["routage_inv"]["choice"] != r["routage"]["choice"]:
                revue.append("contrôle d'ordre : l'orientation change quand les options sont inversées")
            if out["date_promise"] is None:
                dc = ava.date_par_composants(r, arrete)
                if dc:
                    out["date_promise"], out["date_source"] = dc, "composants"
                    out["ecart"] = (dc - d["bascule"]).days
                    out["position"] = "Après la bascule" if dc > d["bascule"] else "Avant la bascule"
                    qc[:] = [x for x in qc if not x.startswith("Promesse datée selon TypeSafe mais date illisible")]
                    if out["position"] == "Après la bascule":
                        qc.append(f"Date promise {out['ecart']} j après la bascule en douteux")
            out["scenarios"] = {k: r[f"scen_{k}"]["noul"] for k in ava.SCENARIOS if f"scen_{k}" in r}
        out["statut_retenu"] = "Revue analyste" if revue else out["statut_ts"]

    if out["segment_douteux"]:
        qc.append(f"Segment « {d['segment']} » douteux au vu du nom")
    if out["intra_groupe"]:
        qc.append("Contrepartie intra-groupe : traitement spécifique")
    out["revue"] = revue
    out["qc"] = qc
    return out


# --------------------------------------------------------------------------- #
# 4. Données : copie anonymisée de démonstration ou fichier réel.
# --------------------------------------------------------------------------- #

COMMENTAIRES_FICTIFS = [
    "Le client a transmis l'ordre de virement ce matin, règlement prévu le 03/10/2026.",
    "Retard de paiement de l'État sur le marché de fournitures. Le mandat est visé au Trésor, encaissement attendu fin octobre.",
    "Client injoignable depuis trois semaines, le magasin est fermé. Visite prévue avec le recouvrement.",
    "Régularisé.",
    "Le salaire n'est plus domicilié chez nous depuis son changement d'employeur. Il promet de payer le 15.",
    "Perte du contrat principal avec le client minier, baisse forte du chiffre d'affaires. Le DG demande un rééchelonnement sur 24 mois.",
    "Promesse de règlement d'ici 10 jours, pas d'autre information.",
    "Le compte est bloqué par un avis à tiers détenteur des impôts. Le client conteste le redressement.",
    "Décès du gérant le mois dernier, la succession est en cours. Les héritiers n'ont pas encore été reçus.",
    "Tension de trésorerie liée au retard de paiement de ses clients distributeurs. Le client apporte un chèque de 40 % fin novembre et le solde en décembre.",
    "RAS",
    "Client à jour, aucun impayé sur le compte.",
    "Les factures présentées pour l'escompte semblent fictives, le débiteur cédé ne les reconnaît pas.",
    "Relance faite par téléphone, le client dit qu'il passera à l'agence la semaine prochaine.",
    "Hospitalisation du client, reprise du travail prévue en novembre. L'épouse s'engage à verser le 20 octobre.",
    "Domiciliation de la subvention annuelle confirmée par la tutelle, virement attendu cette semaine.",
    "Le client a demandé une restructuration avec allongement de 12 mois, dossier transmis au crédit.",
    "Prélèvement rejeté pour un problème de RIB après la migration du compte, régularisation technique en cours.",
    "Le client attend le paiement d'une créance sur une société du groupe. Aucune date donnée.",
    "",
    "",
    "",
    "Paiement prévu le 28/12/2026 après l'encaissement de la campagne arachidière.",
    "Échanges avec le DAF : virement intra-groupe en attente de validation à Lomé, attendu le 08/10/2026.",
]

CLIENTS_FICTIFS = [
    ("SOCIETE NATIONALE DES EAUX DU KAYOR", "Corporate"),
    ("MINISTERE DE L'HYDRAULIQUE ET DE L'ASSAINISSEMENT", "Public Sector"),
    ("COMMUNE DE TAMBA-NORD", "Public Sector"),
    ("AGENCE DE DEVELOPPEMENT DES PORTS FLUVIAUX", "Commercial"),
    ("ECOBANK FOUNDATION SENEGAL", "Corporate"),
    ("EBI SA PARIS", "Corporate"),
    ("DIALLO AMADOU", "Corporate"),
    ("NDIAYE FATOU", "Consumer"),
    ("SOW MOUSTAPHA", "Consumer"),
    ("BOULANGERIE DU PLATEAU SARL", "Commercial"),
    ("SENEGAL LOGISTIQUE ET TRANSIT SA", "Commercial"),
    ("CISSE AWA", "Consumer"),
    ("SOCIETE AFRICAINE DE CIMENT DU SINE SA", "Consumer"),
    ("OFFICE NATIONAL DES SEMENCES", "Commercial"),
    ("GIE DES MARAICHERS DE NIAYES", "Commercial"),
    ("FALL IBRAHIMA", "Consumer"),
    ("BA MARIAMA", "Consumer"),
    ("HOTEL TERANGA BLEU SA", "Corporate"),
    ("SOCIETE DE TRANSPORT URBAIN DE THIES", "Commercial"),
    ("CENTRE HOSPITALIER REGIONAL DE KOLDA", "Commercial"),
    ("GUEYE CHEIKH", "Consumer"),
    ("IMPRIMERIE DU FLEUVE SUARL", "Commercial"),
    ("SECK ADAMA", "Consumer"),
    ("ECOBANK TOGO", "Corporate"),
    ("AGRO-INDUSTRIES DU SALOUM SA", "Corporate"),
    ("FAYE OUMAR", "Consumer"),
    ("KANE KHADY", "Commercial"),
    ("CAISSE DE SECURITE SOCIALE DU PERSONNEL MUNICIPAL", "Commercial"),
    ("QUINCAILLERIE MODERNE DE PIKINE", "Commercial"),
    ("MBAYE ALIOUNE", "Consumer"),
]


def demo(arrete, n=30, graine=29):
    """Copie anonymisée : noms, références et montants fictifs, commentaires fictifs."""
    rnd = random.Random(graine)
    lignes = []
    textes = COMMENTAIRES_FICTIFS + ava.COMMENTAIRES_COMPLEXES  # 24 situations simples + 6 situations complexes
    for i, (client, seg) in enumerate(CLIENTS_FICTIFS[:n]):
        jours = rnd.randint(31, 90)
        enc = rnd.choice([3, 8, 15, 40, 120, 350, 900, 2400]) * 1_000_000 + rnd.randint(0, 999) * 1000
        imp = int(enc * rnd.uniform(0.03, 0.25))
        lignes.append({
            "ref": f"LD{9000001 + i}",
            "client": client,
            "segment": seg,
            "gestionnaire": f"GEST-{rnd.randint(1, 8):02d}",
            "encours": enc,
            "impaye": imp,
            "jours": jours,
            "commentaire": textes[i % len(textes)],
            "statut_fichier": "À faire",
        })
    return lignes


ENTETES = {
    "ref": r"ref|contrat|ld",
    "client": r"client|contrepartie|nom",
    "segment": r"segment",
    "gestionnaire": r"gestionnaire|charge|rm",
    "encours": r"encours",
    "impaye": r"impaye|montant",
    "jours": r"jours|dpd|retard",
    "bascule": r"bascule",
    "commentaire": r"commentaire",
    "statut_fichier": r"statut",
}


def lire_fichier(chemin):
    """Lecture tolérante de la feuille « Plan d'actions » (xlsx/xlsm, ou xlsb avec pyxlsb)."""
    if chemin.suffix.lower() == ".csv":
        with open(chemin, encoding="utf-8") as f:
            lignes = list(csv.reader(f, delimiter=";"))
    elif chemin.suffix.lower() == ".xlsb":
        from pyxlsb import open_workbook          # pip install pyxlsb
        with open_workbook(str(chemin)) as wb:
            nom = next(s for s in wb.sheets if "plan d" in sans_accent(s))
            with wb.get_sheet(nom) as ws:
                lignes = [[c.v for c in r] for r in ws.rows()]
    else:
        import openpyxl
        wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
        nom = next((s for s in wb.sheetnames if "plan d" in sans_accent(s)), wb.sheetnames[0])
        lignes = [list(r) for r in wb[nom].iter_rows(values_only=True)]
    for h, ligne in enumerate(lignes[:25]):
        txt = [sans_accent(str(x or "")) for x in ligne]
        if sum(bool(re.search(p, " ".join(txt))) for p in ENTETES.values()) >= 5:
            break
    col = {}
    for cle, motif in ENTETES.items():
        for j, t in enumerate(txt):
            if j not in col.values() and re.search(motif, t):
                col[cle] = j
                break
    out = []
    for ligne in lignes[h + 1:]:
        if not ligne or ligne[col.get("client", 0)] in (None, ""):
            continue
        g = lambda k, dft=None: ligne[col[k]] if k in col and col[k] < len(ligne) else dft
        out.append({
            "ref": str(g("ref", "")), "client": str(g("client", "")), "segment": str(g("segment", "")),
            "gestionnaire": str(g("gestionnaire", "")), "encours": float(g("encours", 0) or 0),
            "impaye": float(g("impaye", 0) or 0), "jours": int(float(g("jours", 0) or 0)),
            "commentaire": str(g("commentaire", "") or "").strip(),
            "statut_fichier": str(g("statut_fichier", "") or ""),
        })
    return out


# --------------------------------------------------------------------------- #
# 5. Classeur Excel BLUE ECOBANK.
# --------------------------------------------------------------------------- #

MARINE, BLEU, CLAIR, LIME, VERT = "#00415E", "#005C83", "#1A86B3", "#8CC63F", "#4E8A2E"
TEXTE, FOND, LIGNE, ROUGE, OCRE = "#12333F", "#EEF4F7", "#CFE0E7", "#C0392B", "#B67D1C"
PALETTE = [BLEU, LIME, CLAIR, "#0A6F95", OCRE, ROUGE, "#6BA23A", "#7FA7B8", "#D4A13A", "#3E5C6B"]
FEUILLES = ["Tableau de bord", "Plan d'actions", "Retours gestionnaires", "Promesses vs bascules",
            "Contrôle qualité", "Stress narratif", "Méthodologie TypeSafe"]


class Classeur:
    def __init__(self, chemin, feuilles=None):
        self.feuilles = feuilles or FEUILLES
        self.wb = xlsxwriter.Workbook(str(chemin), {"nan_inf_to_errors": True,
                                                     "default_date_format": "dd/mm/yyyy"})
        self.wb.set_properties({"title": "Impayés 30-90j — Couche TypeSafe", "company": "Ecobank Sénégal",
                                "comments": "INTERNAL USE ONLY"})
        f = self.wb.add_format
        base = {"font_name": "Segoe UI", "font_size": 9, "font_color": TEXTE, "valign": "vcenter"}
        self.f = {
            "bandeau": f({**base, "bold": True, "font_size": 16, "font_color": "white", "bg_color": MARINE}),
            "sous": f({**base, "font_size": 9, "font_color": "#DCEAF1", "bg_color": MARINE, "italic": True}),
            "filet": f({"bg_color": LIME}),
            "hdr": f({**base, "bold": True, "font_color": "white", "bg_color": MARINE, "text_wrap": True,
                      "align": "center", "bottom": 2, "bottom_color": LIME, "border_color": LIGNE}),
            "cell": f({**base, "border": 1, "border_color": LIGNE}),
            "wrap": f({**base, "border": 1, "border_color": LIGNE, "text_wrap": True, "valign": "top"}),
            "num": f({**base, "border": 1, "border_color": LIGNE, "num_format": "# ##0", "font_name": "Consolas"}),
            "pct": f({**base, "border": 1, "border_color": LIGNE, "num_format": "0%", "align": "center"}),
            "date": f({**base, "border": 1, "border_color": LIGNE, "num_format": "dd/mm/yyyy", "align": "center"}),
            "int": f({**base, "border": 1, "border_color": LIGNE, "align": "center"}),
            "tile_l": f({**base, "bold": True, "font_size": 8, "font_color": "#3E5C6B", "bg_color": FOND,
                         "left": 5, "left_color": LIME, "text_wrap": True}),
            "tile_v": f({**base, "bold": True, "font_size": 18, "font_color": MARINE, "bg_color": FOND,
                         "left": 5, "left_color": LIME, "num_format": "# ##0"}),
            "nav": f({**base, "bold": True, "font_color": "white", "bg_color": CLAIR, "align": "center",
                      "border": 1, "border_color": "white"}),
            "nav_on": f({**base, "bold": True, "font_color": MARINE, "bg_color": LIME, "align": "center",
                         "border": 1, "border_color": "white"}),
            "h2": f({**base, "bold": True, "font_size": 12, "font_color": MARINE, "bottom": 2, "bottom_color": LIME}),
            "p": f({**base, "text_wrap": True, "valign": "top"}),
            "rouge": f({"font_color": ROUGE, "bold": True}),
            "ocre": f({"bg_color": "#FBEFD9"}),
            "vert": f({"bg_color": "#E6F3D8"}),
            "rose": f({"bg_color": "#F8E1DE"}),
        }

    def entete(self, ws, nom, titre, sous_titre, largeur):
        ws.hide_gridlines(2)
        ws.set_tab_color(BLEU if nom != "_TYPESAFE" else "#7FA7B8")
        ws.set_landscape()
        ws.set_paper(9)
        ws.fit_to_pages(1, 0)
        ws.set_footer("&L&8ECOBANK SÉNÉGAL · Credit Risk · INTERNAL USE ONLY&R&8&P / &N")
        ws.set_row(0, 30)
        ws.merge_range(0, 0, 0, largeur - 1, titre, self.f["bandeau"])
        ws.set_row(1, 18)
        ws.merge_range(1, 0, 1, largeur - 1, sous_titre, self.f["sous"])
        ws.set_row(2, 4)
        for c in range(largeur):
            ws.write_blank(2, c, None, self.f["filet"])
        # Barre de navigation : un bouton par feuille, feuille active en surbrillance.
        ws.set_row(3, 22)
        F = self.feuilles
        par = max(1, largeur // len(F))
        for i, cible in enumerate(F):
            c0 = i * par
            c1 = largeur - 1 if i == len(F) - 1 else c0 + par - 1
            fmt = self.f["nav_on"] if cible == nom else self.f["nav"]
            lien = f"internal:'{cible}'!A1"
            if c1 > c0:
                ws.merge_range(3, c0, 3, c1, "", fmt)
            ws.write_url(3, c0, lien, fmt, cible)

    def tableau(self, ws, ligne0, colonnes, lignes, nom_table):
        """colonnes = [(titre, largeur, clé ou fonction, format)]"""
        ws.set_row(ligne0, 36)   # en-têtes sur deux lignes : plus de titre coupé
        for j, (titre, larg, _, _) in enumerate(colonnes):
            ws.set_column(j, j, larg)
        data = []
        for r in lignes:
            data.append([(k(r) if callable(k) else r.get(k)) for _, _, k, _ in colonnes])
        ws.add_table(ligne0, 0, ligne0 + max(1, len(data)), len(colonnes) - 1, {
            "name": nom_table, "style": "Table Style Light 9", "autofilter": True,
            "data": data or [[None] * len(colonnes)],
            "columns": [{"header": t, "header_format": self.f["hdr"], "format": self.f[fm]}
                        for t, _, _, fm in colonnes],
        })
        for i, r in enumerate(data, 1):
            hauteur = max(15, 13 * max((len(str(v)) // max(8, int(colonnes[j][1] * 1.1)) + 1)
                                       for j, v in enumerate(r) if v is not None and colonnes[j][3] == "wrap") if
                          any(colonnes[j][3] == "wrap" for j in range(len(r))) else 15)
            ws.set_row(ligne0 + i, min(hauteur, 90))
        ws.freeze_panes(ligne0 + 1, 2)
        return ligne0 + len(data)

    def close(self):
        self.wb.close()


def ajouter_serie(ch, nom, cat_ref, val_ref, valeurs, couleurs, donut=False, masquer_sous=0.04):
    s = {"name": nom, "categories": cat_ref, "values": val_ref,
         "points": [{"fill": {"color": couleurs[i % len(couleurs)]}, "border": {"color": "white"}}
                    for i in range(len(valeurs))]}
    total = sum(valeurs) or 1
    if donut:
        # Étiquettes masquées sous 4 % : plus de chevauchement dans l'anneau.
        s["data_labels"] = {"percentage": True, "position": "best_fit",
                            "leader_lines": True, "font": {"name": "Segoe UI", "size": 8, "color": TEXTE},
                            "custom": [{"delete": True} if v / total < masquer_sous else None for v in valeurs]}
    else:
        s["data_labels"] = {"value": True, "font": {"name": "Segoe UI", "size": 8, "color": TEXTE}}
        s["gap"] = 60
    ch.add_series(s)


def style_graphique(ch, titre, legende=True):
    ch.set_title({"name": titre, "name_font": {"name": "Segoe UI", "size": 11, "bold": True, "color": MARINE},
                  "overlay": False})
    ch.set_legend({"position": "right", "font": {"name": "Segoe UI", "size": 8}} if legende else {"none": True})
    ch.set_chartarea({"border": {"color": LIGNE}, "fill": {"color": "white"}})
    ch.set_plotarea({"fill": {"color": "white"}})


def evaluer_portefeuille(pf, essais=5):
    """Un appel « portefeuille » : importance de chaque constat, phrase de lecture, contreparties liées."""
    etat = {"insights": [{"id": x["id"], "titre": x["titre"], "texte": x["texte"]} for x in pf.get("insights", [])],
            "paires": [{"id": p["id"], "a": p["a"], "b": p["b"]} for p in pf.get("paires", [])]}
    corps = json.dumps({"model": MODELE, "state": etat, "questions": ava.questions_portefeuille(pf)}).encode()
    entetes = {"Content-Type": "application/json"}
    if os.environ.get("TYPESAFE_API_KEY"):
        entetes["Authorization"] = f"Bearer {os.environ['TYPESAFE_API_KEY']}"
    for essai in range(essais):
        try:
            with urllib.request.urlopen(urllib.request.Request(API_URL, data=corps, headers=entetes), timeout=120) as rep:
                js = json.load(rep)
                return {"modele": MODELE, "version": js.get("model"), "usage": js.get("usage"),
                        "horodatage": datetime.now().isoformat(timespec="seconds"), "answers": js["answers"]}
        except (urllib.error.HTTPError, urllib.error.URLError):
            if essai < essais - 1:
                time.sleep(2 ** essai)
                continue
            raise


def ecrire_classeur(resultats, bruts, arrete, chemin, portefeuille=None):
    C = Classeur(chemin)
    wb, f = C.wb, C.f
    n = len(resultats)
    avec_retour = [r for r in resultats if r["commentaire"]]
    revue = [r for r in resultats if r["revue"]]
    apres = [r for r in resultats if r["position"] == "Après la bascule"]
    impaye = sum(r["impaye"] for r in resultats)

    # ---------- Tableau de bord ----------
    ws = wb.add_worksheet("Tableau de bord")
    LARG = 12
    ws.set_column(0, LARG - 1, 13.5)   # 12 colonnes uniformes ≈ 1 220 px : les graphiques tiennent dedans
    C.entete(ws, "Tableau de bord", "IMPAYÉS 30-90J · RETOURS GESTIONNAIRES · COUCHE TYPESAFE",
             f"Arrêté du {arrete:%d/%m/%Y} · copie anonymisée de démonstration · chiffres calculés par le code, "
             f"jugements TypeSafe ({MODELE})", LARG)
    tuiles = [
        ("Dossiers", n), ("Impayé total (XOF)", impaye), ("Avec commentaire", len(avec_retour)),
        ("Revue analyste", len(revue)), ("Promesse après bascule", len(apres)),
        ("Public / parapublic", sum(r["nature"] in ("Publique", "Parapublique") for r in resultats)),
    ]
    ws.set_row(5, 26)
    ws.set_row(6, 30)
    for i, (lib, v) in enumerate(tuiles):
        ws.merge_range(5, 2 * i, 5, 2 * i + 1, lib.upper(), f["tile_l"])
        ws.merge_range(6, 2 * i, 6, 2 * i + 1, v, f["tile_v"])

    # Données des graphiques (feuille masquée de calcul, pas de formules volatiles).
    wd = wb.add_worksheet("_graph")
    wd.hide()

    def bloc(col, titre, paires):
        wd.write(0, col, titre)
        for i, (k, v) in enumerate(paires, 1):
            wd.write(i, col, k)
            wd.write(i, col + 1, v)
        m = len(paires)
        return (["_graph", 1, col, m, col], ["_graph", 1, col + 1, m, col + 1], [v for _, v in paires])

    def compte(cle, ordre=None):
        d = {}
        for r in resultats:
            d[r[cle]] = d.get(r[cle], 0) + 1
        cles = ordre or sorted(d, key=lambda k: -d[k])
        return [(k, d.get(k, 0)) for k in cles if d.get(k, 0)]

    def somme(cle, ordre):
        d = {}
        for r in resultats:
            d[r[cle]] = d.get(r[cle], 0) + r["impaye"] / 1e6
        return [(k, round(d.get(k, 0), 1)) for k in ordre if d.get(k)]

    motifs = bloc(0, "Motif", compte("motif"))
    classes = bloc(3, "Classe", somme("classe", ["Sensible 31-60j", "Sensible 61-90j", "Douteux"]))
    statuts = bloc(6, "Statut", compte("statut_retenu"))
    position = bloc(9, "Position", compte("position", ["Avant la bascule", "Après la bascule", "Sans date", "Sans retour"]))
    natures = bloc(12, "Nature", compte("nature", ["Publique", "Parapublique", "Privée", "Particulier"]))
    creds = bloc(15, "Crédibilité", [(LIB_CRED[k], sum(r["credibilite"] == k for r in avec_retour))
                                     for k in range(5)])

    W, H = 600, 300          # deux graphiques par rangée, 1 200 px : rien n'est coupé à droite
    def poser(ch, ligne, col, x=0):
        ch.set_size({"width": W, "height": H})
        ws.insert_chart(ligne, col, ch, {"x_offset": x + 6, "y_offset": 6, "object_position": 1})

    ch = wb.add_chart({"type": "doughnut"})
    ajouter_serie(ch, "Dossiers", motifs[0], motifs[1], motifs[2], PALETTE, donut=True)
    ch.set_hole_size(55)
    style_graphique(ch, "Motifs des impayés")
    poser(ch, 8, 0)

    ch = wb.add_chart({"type": "bar"})
    ajouter_serie(ch, "Impayé (M XOF)", classes[0], classes[1], classes[2], [CLAIR, OCRE, ROUGE])
    style_graphique(ch, "Impayé par classe (M XOF)", legende=False)
    ch.set_x_axis({"num_font": {"size": 8}, "major_gridlines": {"visible": True, "line": {"color": FOND}}})
    ch.set_y_axis({"num_font": {"size": 9}, "reverse": True})
    poser(ch, 8, 6)

    ch = wb.add_chart({"type": "bar"})
    ajouter_serie(ch, "Dossiers", statuts[0], statuts[1], statuts[2], PALETTE)
    style_graphique(ch, "Statut du suivi proposé", legende=False)
    ch.set_y_axis({"num_font": {"size": 8}, "reverse": True})
    ch.set_x_axis({"num_font": {"size": 8}, "major_gridlines": {"visible": False}})
    poser(ch, 24, 0)

    ch = wb.add_chart({"type": "column"})
    ajouter_serie(ch, "Dossiers", position[0], position[1], position[2], [LIME, ROUGE, OCRE, "#7FA7B8"])
    style_graphique(ch, "Promesses vs bascule", legende=False)
    ch.set_x_axis({"num_font": {"size": 8}})
    ch.set_y_axis({"major_gridlines": {"visible": True, "line": {"color": FOND}}})
    poser(ch, 24, 6)

    ch = wb.add_chart({"type": "doughnut"})
    ajouter_serie(ch, "Dossiers", natures[0], natures[1], natures[2], [MARINE, CLAIR, LIME, OCRE], donut=True)
    ch.set_hole_size(55)
    style_graphique(ch, "Nature des contreparties")
    poser(ch, 40, 0)

    ch = wb.add_chart({"type": "column"})
    ajouter_serie(ch, "Dossiers", creds[0], creds[1], creds[2], [ROUGE, OCRE, "#D4A13A", CLAIR, LIME])
    style_graphique(ch, "Crédibilité des promesses", legende=False)
    ch.set_x_axis({"num_font": {"size": 8}})
    poser(ch, 40, 6)

    lecture = [
        f"{len(avec_retour)} dossiers sur {n} ont un commentaire gestionnaire ; {n - len(avec_retour)} restent sans retour.",
        f"{len(revue)} dossiers partent en revue analyste (confiance faible ou incohérence) : aucune décision automatique.",
        f"{len(apres)} promesses de règlement tombent après la date de bascule en douteux.",
        f"{sum(r['intra_groupe'] for r in resultats)} contrepartie(s) intra-groupe et "
        f"{sum(r['segment_douteux'] for r in resultats)} segment(s) douteux signalés au vu des noms.",
    ]
    ws.merge_range(56, 0, 56, LARG - 1, "Lecture", f["h2"])
    for i, t in enumerate(lecture):
        ws.set_row(57 + i, 18)
        ws.merge_range(57 + i, 0, 57 + i, LARG - 1, "• " + t, f["p"])

    # ---------- Plan d'actions enrichi ----------
    ws = wb.add_worksheet("Plan d'actions")
    cols = [
        ("Réf LD", 12, "ref", "cell"), ("Client", 30, "client", "wrap"), ("Segment", 12, "segment", "cell"),
        ("Gestion-naire", 10, "gestionnaire", "cell"), ("Encours (XOF)", 15, "encours", "num"),
        ("Impayé (XOF)", 14, "impaye", "num"), ("Jours d'impayés", 9, "jours", "int"),
        ("Classe", 14, "classe", "cell"), ("Date de bascule", 11, "bascule", "date"),
        ("Commentaire du gestionnaire", 46, "commentaire", "wrap"),
        ("Statut du suivi (fichier)", 12, "statut_fichier", "cell"),
        ("Statut proposé TypeSafe", 16, lambda r: r["statut_ts"] or "—", "cell"),
        ("Statut retenu", 16, "statut_retenu", "cell"),
        ("Motif (TypeSafe)", 18, "motif", "cell"),
        ("Revue analyste", 34, lambda r: " ; ".join(r["revue"]) or "", "wrap"),
    ]
    C.entete(ws, "Plan d'actions", "PLAN D'ACTIONS · IMPAYÉS 30-90J · ENRICHI TYPESAFE",
             "Colonnes TypeSafe à droite ; le statut retenu vaut « Revue analyste » quand la confiance est faible.",
             len(cols))
    fin = C.tableau(ws, 5, cols, resultats, "PlanActions")
    ws.conditional_format(6, 12, fin, 12, {"type": "cell", "criteria": "==", "value": '"Revue analyste"',
                                            "format": f["rose"]})
    ws.conditional_format(6, 5, fin, 5, {"type": "data_bar", "bar_color": CLAIR, "bar_solid": True})
    ws.data_validation(6, 12, fin, 12, {"validate": "list", "source": list(LIB_STATUT.values()) + ["Revue analyste"]})

    # ---------- Retours gestionnaires ----------
    ws = wb.add_worksheet("Retours gestionnaires")
    cols = [
        ("Réf LD", 12, "ref", "cell"), ("Client", 28, "client", "wrap"), ("Gestion-naire", 10, "gestionnaire", "cell"),
        ("Commentaire du gestionnaire", 48, "commentaire", "wrap"),
        ("Motif", 18, "motif", "cell"), ("Confiance motif", 10, "conf_motif", "pct"),
        ("Promesse datée (proba.)", 10, "p_promesse", "pct"),
        ("Crédibilité", 11, lambda r: LIB_CRED[r["credibilite"]] if r["credibilite"] is not None else "—", "cell"),
        ("Confiance crédibilité", 11, "conf_cred", "pct"),
        ("Incohérence (proba.)", 11, "p_incoherence", "pct"),
        ("Statut proposé", 16, "statut_ts", "cell"), ("Confiance statut", 10, "conf_statut", "pct"),
        ("Statut retenu", 16, "statut_retenu", "cell"),
        ("Urgence (0-4)", 9, lambda r: r.get("urgence"), "int"), ("Orientation", 20, lambda r: r.get("routage") or "—", "cell"),
        ("Famille de motif", 18, lambda r: r.get("famille") or "—", "cell"), ("Action recommandée", 22, lambda r: r.get("action") or "—", "cell"),
        ("Stabilité", 9, lambda r: r.get("stabilite"), "pct"),
        ("Scénarios exposés", 30, lambda r: ", ".join(k.replace("_", " ") for k, v in (r.get("scenarios") or {}).items() if v >= SEUIL_SIGNAL) or "—", "wrap"),
        ("Durabilité (proba.)", 10, lambda r: r.get("p_durable"), "pct"), ("Plan cohérent (proba.)", 10, lambda r: r.get("p_coherence_plan"), "pct"),
    ]
    C.entete(ws, "Retours gestionnaires", "RETOURS GESTIONNAIRES · LECTURE TYPESAFE DES COMMENTAIRES",
             "Une ligne par dossier commenté. Probabilités et confiances brutes conservées dans _TYPESAFE.", len(cols))
    fin = C.tableau(ws, 5, cols, avec_retour, "Retours")
    for c in (5, 8, 11):
        ws.conditional_format(6, c, fin, c, {"type": "3_color_scale", "min_color": "#F8E1DE",
                                              "mid_color": "#FBEFD9", "max_color": "#E6F3D8"})
    ws.conditional_format(6, 9, fin, 9, {"type": "cell", "criteria": ">=", "value": SEUIL_SIGNAL, "format": f["rose"]})

    # ---------- Promesses vs bascules ----------
    ws = wb.add_worksheet("Promesses vs bascules")
    prom = sorted([r for r in avec_retour if r["date_promise"] or r["p_promesse"] >= SEUIL_SIGNAL],
                  key=lambda r: (r["ecart"] is None, -(r["ecart"] or 0)))
    cols = [
        ("Réf LD", 12, "ref", "cell"), ("Client", 28, "client", "wrap"), ("Impayé (XOF)", 14, "impaye", "num"),
        ("Jours d'impayés", 9, "jours", "int"), ("Date de bascule", 11, "bascule", "date"),
        ("Date promise (lue par le code)", 13, "date_promise", "date"),
        ("Écart (jours)", 9, "ecart", "int"), ("Position", 16, "position", "cell"),
        ("Crédibilité", 11, lambda r: LIB_CRED[r["credibilite"]], "cell"),
        ("Commentaire du gestionnaire", 50, "commentaire", "wrap"),
    ]
    C.entete(ws, "Promesses vs bascules", "PROMESSES DE RÈGLEMENT COMPARÉES À LA DATE DE BASCULE",
             "Bascule = date à laquelle le dossier atteint 91 jours (calcul du code). Écart > 0 : la promesse arrive trop tard.",
             len(cols))
    fin = C.tableau(ws, 5, cols, prom, "Promesses")
    ws.conditional_format(6, 7, fin, 7, {"type": "cell", "criteria": "==", "value": '"Après la bascule"', "format": f["rose"]})
    ws.conditional_format(6, 7, fin, 7, {"type": "cell", "criteria": "==", "value": '"Avant la bascule"', "format": f["vert"]})
    ws.conditional_format(6, 6, fin, 6, {"type": "cell", "criteria": ">", "value": 0, "format": f["rouge"]})

    # ---------- Contrôle qualité ----------
    ws = wb.add_worksheet("Contrôle qualité")
    cq = [dict(r, point=p) for r in resultats for p in (r["qc"] + [f"Revue : {x}" for x in r["revue"]])]
    cols = [
        ("Réf LD", 12, "ref", "cell"), ("Client", 30, "client", "wrap"), ("Segment", 13, "segment", "cell"),
        ("Gestion-naire", 10, "gestionnaire", "cell"), ("Nature (TypeSafe)", 13, "nature", "cell"),
        ("Intra-groupe", 9, lambda r: "Oui" if r["intra_groupe"] else "Non", "int"),
        ("Segment douteux", 9, lambda r: "Oui" if r["segment_douteux"] else "Non", "int"),
        ("Point de contrôle", 54, "point", "wrap"),
    ]
    C.entete(ws, "Contrôle qualité", "CONTRÔLE QUALITÉ DES RETOURS ET DES NOMS DE CLIENTS",
             "Commentaires vides ou trop courts, dates illisibles, incohérences, segments douteux, intra-groupe, revues.",
             len(cols))
    fin = C.tableau(ws, 5, cols, cq, "Controle")
    ws.conditional_format(6, 5, fin, 6, {"type": "cell", "criteria": "==", "value": '"Oui"', "format": f["ocre"]})

    # ---------- Stress narratif (exposition jugée, montants agrégés en code) ----------
    ws = wb.add_worksheet("Stress narratif")
    sc_rows = []
    for k, lib in ava.SCENARIOS.items():
        exp_ = [r for r in resultats if (r.get("scenarios") or {}).get(k, 0) >= SEUIL_SIGNAL]
        sc_rows.append({"sc": lib, "n": len(exp_), "enc": sum(r["encours"] for r in exp_), "imp": sum(r["impaye"] for r in exp_),
                        "cl": ", ".join(r["client"] for r in exp_[:6]) + (" …" if len(exp_) > 6 else "")})
    cols = [("Scénario", 46, "sc", "wrap"), ("Dossiers exposés", 10, "n", "int"), ("Encours exposé (XOF)", 18, "enc", "num"),
            ("Impayé exposé (XOF)", 16, "imp", "num"), ("Principaux dossiers", 60, "cl", "wrap")]
    C.entete(ws, "Stress narratif", "STRESS NARRATIF · EXPOSITION DES DOSSIERS AUX SCÉNARIOS MACRO",
             "Exposition jugée par le modèle (probabilité ≥ 0,6) ; montants agrégés par le code.", len(cols))
    fin = C.tableau(ws, 5, cols, sc_rows, "StressNarratif")
    ws.conditional_format(6, 2, fin, 2, {"type": "data_bar", "bar_color": ROUGE, "bar_solid": True})

    # ---------- Méthodologie ----------
    ws = wb.add_worksheet("Méthodologie TypeSafe")
    ws.set_column(0, 0, 30)
    ws.set_column(1, 1, 110)
    C.entete(ws, "Méthodologie TypeSafe", "MÉTHODOLOGIE ET PISTE D'AUDIT",
             f"Modèle {MODELE} · {API_URL} · seuil de signal {SEUIL_SIGNAL} · seuil de confiance {SEUIL_CONFIANCE}", 2)
    lignes = [
        ("Partage des rôles", "Le code calcule tout ce qui est réglementaire ou chiffré (BCEAO, IFRS9, ACTE7) : classe, jours, "
         "date de bascule, montants, date promise lue dans le texte, écart. TypeSafe ne rend que des jugements : motif, "
         "promesse datée, crédibilité, incohérence, statut proposé, nature de l'entité, intra-groupe, segment douteux. "
         "Il ne produit jamais de montant ni de ratio."),
        ("Date de bascule", f"Arrêté + ({SEUIL_DOUTEUX + 1} − jours d'impayés) : date à laquelle le dossier devient douteux "
         "si rien n'est réglé."),
        ("Date promise", "Lue par le code (jj/mm/aaaa, « le 15 », « fin octobre », « sous 10 jours », « cette semaine »…). "
         "Si TypeSafe voit une promesse datée que le code ne sait pas lire, le contrôle qualité le signale."),
        ("Revue analyste", f"Confiance < {SEUIL_CONFIANCE} sur le motif, le statut, la crédibilité ou la nature, ou "
         f"incohérence ≥ {SEUIL_SIGNAL} : le statut retenu devient « Revue analyste », sans décision automatique."),
        ("Statut retenu", "Statut proposé par TypeSafe quand la confiance suffit, sinon « Revue analyste ». Le gestionnaire "
         "et l'analyste restent seuls décisionnaires (liste déroulante dans Plan d'actions)."),
        ("Réponses brutes", "Feuille masquée _TYPESAFE : probabilités, confiance, modèle et horodatage par dossier. "
         "Changer les seuils dans impayes_typesafe.py puis relancer avec --depuis-cache ne coûte aucun appel."),
        ("Offline", "Les postes bancaires n'ont pas Internet. L'enrichissement tourne sur un poste connecté ou un relais "
         "serveur ; l'app HTML relit _TYPESAFE et fonctionne sans elle. Aucune clé API n'est écrite dans un fichier."),
        ("Données", "Démonstration sur une copie anonymisée : noms, références et montants fictifs, commentaires fictifs. "
         "Les données réelles n'iront à l'API qu'après validation de la Conformité."),
        ("Capacités avancées", "Re-classement des dossiers par urgence (Score), routage vers le bon traitement (Choice), contrôle d'auto-cohérence : le motif (Choice) et sa durabilité (Noul) sont posés séparément ; une divergence envoie le dossier en revue analyste. Cohérence entre le plan d'action APEX et la cause décrite (Noul). Toutes les questions partent en un seul appel (éventail spéculatif)."),
        ("Découvertes API", "Version exacte du modèle et jetons consommés conservés (piste d'audit). Trois types de questions seulement (choice, score, noul). Français aussi fiable que l'anglais sur ce domaine (21/21, confiance 0,916 / 0,914). Ordre des options inversé : 21/21 identiques, contrôle conservé sur motif et routage. Sel dans l'état : ±0,03 de variation, utilisé pour mesurer la stabilité des dossiers sensibles (3 tirages)."),
        ("Stress narratif", "Huit scénarios macro (paiements de l'État, campagne agricole, hydrocarbures, taux BCEAO, crise sanitaire, perte d'un donneur d'ordre, prix à l'importation, gouvernance) : le modèle juge l'exposition directe de chaque dossier (Noul) ; le code agrège les encours exposés et applique le stress de Risk Outlook."),
        ("Apprentissage supervisé", "Le modèle n'est pas réentraîné (l'API ne le permet pas) : ce sont ses consignes et nos seuils "
         "qui apprennent. Les analystes valident ou corrigent les jugements dans a_valider.xlsx (--modele-apprentissage) ; "
         "les réponses validées (--apprentissage) sont jointes en exemples aux questions motif, orientation, statut et action "
         "(2 par option, 12 au plus, noms masqués), et le seuil de confiance de chaque question est recalibré dès 10 cas : "
         "plus petit seuil (plancher 0,40) donnant 90 % d'accord avec les analystes. État : " + app.resume()),
        ("Questions posées", "Commentaire : " + " · ".join(Q_COMMENTAIRE) + ". Nom du client : " + " · ".join(Q_NOM) + "."),
    ]
    ws.write_row(5, 0, ["Élément", "Règle"], f["hdr"])
    for i, (k, v) in enumerate(lignes, 6):
        ws.set_row(i, 15 * (len(v) // 120 + 1) + 6)
        ws.write(i, 0, k, f["wrap"])
        ws.write(i, 1, v, f["wrap"])

    # ---------- _TYPESAFE (masquée) ----------
    ws = wb.add_worksheet("_TYPESAFE")
    C.entete(ws, "_TYPESAFE", "_TYPESAFE · RÉPONSES BRUTES (PISTE D'AUDIT)",
             "Lue par l'app HTML offline au rechargement. Ne pas modifier.", 6)
    ws.write_row(5, 0, ["ref", "question", "type", "reponse", "confiance", "probabilites_json", "libelle"], f["hdr"])
    versions = sorted({b.get("version") for b in bruts.values() if isinstance(b, dict) and b.get("version")})
    jetons = sum((b.get("usage") or {}).get("input_tokens", 0) for b in bruts.values() if isinstance(b, dict))
    ws.write_row(4, 0, ["modele", ", ".join(versions) or MODELE, "arrete", arrete.isoformat(), "genere", datetime.now().isoformat(timespec="seconds"), f"jetons d'entrée : {jetons}"])
    ws.set_column(0, 4, 16)
    ws.set_column(5, 5, 90)
    i = 6
    for r in resultats:
        b = bruts[cle(r)]
        for q, a in b["answers"].items():
            rep = a.get("choice", a.get("score", a.get("noul")))
            ws.write_row(i, 0, [r["ref"], q, a["type"], rep, a.get("confidence"),
                                json.dumps(a.get("probabilities", {}), ensure_ascii=False)])
            i += 1
        if r.get("stabilite") is not None:
            ws.write_row(i, 0, [r["ref"], "stabilite", "code", "stable" if r["stable"] else "instable", r["stabilite"], "", "3 rééchantillonnages (sel dans l'état)"])
            i += 1
    if portefeuille and portefeuille[1]:
        pf, b = portefeuille
        lib = {**{f"importance_{x['id']}": x["titre"] for x in pf.get("insights", [])},
               **{f"groupe_{p['id']}": f"{p['a']} ↔ {p['b']}" for p in pf.get("paires", [])},
               **{f"meme_nom_{p['id']}": f"{p['a']} ↔ {p['b']}" for p in pf.get("paires", [])}}
        for q, a in b["answers"].items():
            rep = a.get("choice", a.get("score", a.get("noul")))
            ws.write_row(i, 0, ["__PORTEFEUILLE__", q, a["type"], rep, a.get("confidence"), json.dumps(a.get("probabilities", {}), ensure_ascii=False), lib.get(q, "")])
            i += 1
    ws.hide()
    wb.worksheets()[0].activate()
    C.close()


# --------------------------------------------------------------------------- #

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--demo", action="store_true", help="copie anonymisée et commentaires fictifs")
    ap.add_argument("--fichier", type=Path, help="fichier retourné par les gestionnaires (.xlsx, .xlsb, .csv)")
    ap.add_argument("--lot", type=Path, help="lot JSON exporté par APEX (bouton « Lot TypeSafe »)")
    ap.add_argument("--anonymiser", action="store_true",
                    help="masque les noms clients avant l'appel (sans accord Conformité) ; les jugements sur les noms sont alors sans objet")
    ap.add_argument("--arrete", default="2026-09-29")
    ap.add_argument("--depuis-cache", action="store_true")
    ap.add_argument("--modele-apprentissage", action="store_true",
                    help="écrit sortie/a_valider.xlsx : jugements proposés à valider par les analystes")
    ap.add_argument("--apprentissage", type=Path, help="intègre un fichier a_valider.xlsx rempli par les analystes")
    ap.add_argument("--sortie", type=Path, default=SORTIE / "Impayes_30-90j_TypeSafe_demo.xlsx")
    a = ap.parse_args()
    arrete = date.fromisoformat(a.arrete)
    if a.apprentissage:
        ajout, maj = app.integrer(a.apprentissage, LIBELLES_APPRENTISSAGE,
                                  lambda c, client: ava.masquer(c, {client} if client else set(), "[client]"))
        print(f"Apprentissage : {ajout} réponse(s) ajoutée(s), {maj} mise(s) à jour → {app.MEMOIRE}")
        print(app.resume())
        return

    if a.demo:
        dossiers = demo(arrete)
        with open(DEMO_CSV, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Réf LD", "Client", "Segment", "Gestionnaire", "Encours", "Montant impayé",
                        "Jours d'impayés", "Commentaire du gestionnaire", "Statut du suivi"])
            for d in dossiers:
                w.writerow([d["ref"], d["client"], d["segment"], d["gestionnaire"], d["encours"], d["impaye"],
                            d["jours"], d["commentaire"], d["statut_fichier"]])
    elif a.fichier:
        dossiers = lire_fichier(a.fichier)
    elif a.lot:
        lot = json.loads(a.lot.read_text(encoding="utf-8"))
        if lot.get("arrete") and a.arrete == "2026-09-29":
            arrete = date.fromisoformat(lot["arrete"])
        dossiers = [{"ref": str(d["ref"]), "client": d.get("client") or "", "segment": d.get("segment") or "",
                     "gestionnaire": d.get("gestionnaire") or "", "encours": float(d.get("encours") or 0),
                     "impaye": float(d.get("impaye") or 0), "jours": int(d.get("jours") or 0),
                     "commentaire": (d.get("commentaire") or "").strip(), "statut_fichier": d.get("statut_fichier") or "",
                     "plan": (d.get("plan") or "").strip(), "secteur": d.get("secteur"), "garantie": d.get("garantie"),
                     "taille": d.get("taille"), "trajectoire": d.get("trajectoire"),
                     "bascule_apex": d.get("bascule")} for d in lot["dossiers"]]
        noms_reels = sorted({d["client"] for d in dossiers if d["client"]})
        if a.anonymiser:
            # un alias stable par client ; noms clients et gestionnaires masqués aussi dans le texte du plan
            alias = {n: f"CLIENT-{i:04d}" for i, n in enumerate(noms_reels, 1)}
            gest = {d["gestionnaire"] for d in dossiers if d["gestionnaire"]}
            for d in dossiers:
                d["plan"] = ava.masquer(ava.masquer(d["plan"], alias, alias), gest, "[gestionnaire]")
                d["commentaire"] = ava.masquer(ava.masquer(d["commentaire"], alias, alias), gest, "[gestionnaire]")
                d["client"] = alias.get(d["client"], d["client"])
        if a.sortie == SORTIE / "Impayes_30-90j_TypeSafe_demo.xlsx":
            a.sortie = SORTIE / f"Impayes_30-90j_TypeSafe_{arrete:%Y-%m-%d}.xlsx"
    else:
        ap.error("--demo, --fichier ou --lot")

    for d in dossiers:
        d["classe"] = classe_impaye(d["jours"])
        # la date de bascule d'APEX fait foi quand elle est fournie (seuil 90 j ou 180 j secteur public)
        d["bascule"] = date.fromisoformat(d["bascule_apex"]) if d.get("bascule_apex") else date_bascule(arrete, d["jours"])

    SORTIE.mkdir(exist_ok=True)
    cache = json.loads(CACHE.read_text()) if CACHE.exists() else {}
    if not a.depuis_cache:
        a_faire = [d for d in dossiers if cle(d) not in cache]
        print(f"TypeSafe : {len(a_faire)} requête(s) ({MODELE})")
        with ThreadPoolExecutor(8) as ex:
            for d, rep in zip(a_faire, ex.map(evaluer, a_faire)):
                cache[cle(d)] = rep
        CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=1))
    manquants = [d["ref"] for d in dossiers if cle(d) not in cache]
    if manquants:
        raise SystemExit(f"Réponses absentes du cache : {manquants[:5]}")

    resultats = [analyser(d, cache[cle(d)], arrete) for d in dossiers]
    # stabilité : 3 rééchantillonnages (sel dans l'état) pour les dossiers urgents ou en revue
    sensibles = [d for d, r in zip(dossiers, resultats) if d["commentaire"] and ((r.get("urgence") or 0) >= 3 or r["revue"])]
    if sensibles and not a.depuis_cache:
        taches = [dict(d, _sel=1, _k=k) for d in sensibles for k in range(3) if f"{cle(d)}|sel{k}" not in cache]
        with ThreadPoolExecutor(8) as ex:
            for t, rep in zip(taches, ex.map(evaluer, taches)):
                cache[f"{cle(t)}|sel{t['_k']}"] = rep
        CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=1))
    for d, r in zip(dossiers, resultats):
        tirs = [cache.get(f"{cle(d)}|sel{k}") for k in range(3)]
        tirs = [t for t in tirs if t]
        if tirs:
            base = cache[cle(d)]["answers"]
            ch = [t["answers"]["motif"]["choice"] for t in tirs] + [t["answers"]["routage"]["choice"] for t in tirs]
            ref = [base["motif"]["choice"]] * len(tirs) + [base["routage"]["choice"]] * len(tirs)
            pm = [t["answers"]["motif"]["probabilities"][base["motif"]["choice"]] for t in tirs] + [base["motif"]["probabilities"][base["motif"]["choice"]]]
            r["stabilite"] = round(1 - (max(pm) - min(pm)), 3)
            r["stable"] = ch == ref
            if not r["stable"]:
                r["revue"].append("instabilité : le jugement change d'un tirage à l'autre")
                r["statut_retenu"] = "Revue analyste"
    # appel portefeuille : phrase de lecture du Comité et contreparties potentiellement liées
    pf = (lot.get("portefeuille") if a.lot else None) or {}
    if a.anonymiser and pf:
        # sans accord Conformité : noms masqués dans les constats, jugements sur les paires de noms sans objet
        cites = set(pf.get("noms", [])) | set(noms_reels) | {d["gestionnaire"] for d in dossiers if d["gestionnaire"]} \
            | {p[k] for p in pf.get("paires", []) for k in ("a", "b")}
        pf = {"insights": [dict(x, titre=ava.masquer(x["titre"], cites), texte=ava.masquer(x["texte"], cites))
                           for x in pf.get("insights", [])], "paires": []}
        print("Anonymisation : noms masqués dans les constats, paires de noms retirées de l'appel portefeuille")
    portefeuille = None
    if pf.get("insights") or pf.get("paires"):
        kpf = "__PORTEFEUILLE__|" + __import__("hashlib").sha1(json.dumps(pf, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:12]
        if kpf not in cache and not a.depuis_cache:
            cache[kpf] = evaluer_portefeuille(pf)
            CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=1))
        portefeuille = (pf, cache.get(kpf))
    ecrire_classeur(resultats, cache, arrete, a.sortie, portefeuille)
    if a.modele_apprentissage:
        n = app.ecrire_a_valider(dossiers, resultats, cache, cle, LIBELLES_APPRENTISSAGE, SORTIE / "a_valider.xlsx",
                                 lambda c, d: c)
        print(f"Apprentissage : {n} jugement(s) à valider → {SORTIE / 'a_valider.xlsx'}")
    (SORTIE / "impayes_resultats.json").write_text(json.dumps(resultats, ensure_ascii=False, indent=1, default=str))

    for r in resultats:
        print(f"{r['ref']} {r['client'][:32]:32} {r['motif'][:22]:22} {r['statut_retenu'][:22]:22} "
              f"{r['position']:17} {r['nature']:12}{' INTRA' if r['intra_groupe'] else ''}"
              f"{' SEG?' if r['segment_douteux'] else ''}")
    print(f"→ {a.sortie}")


if __name__ == "__main__":
    main()
