#!/usr/bin/env python3
"""Exports propres et contrôlés : Excel mis en forme, CSV « Excel français », rapprochement, nettoyage des métadonnées.

    py export_pro.py excel   source.xlsx|csv  --sortie rapport.xlsx [--feuille F] [--titre "…"] [--classe "…"]
    py export_pro.py csv     source.xlsx|csv  --sortie export.csv   [--sep ";"] [--decimale ","]
    py export_pro.py controle source.xlsx|csv  export.xlsx|csv [--cle "Compte"]     # code 1 si écart
    py export_pro.py metadonnees fichier.docx|xlsx|pptx [--nettoyer]               # auteur, société, commentaires…

Rien ne sort de la machine. La source n'est jamais modifiée. Chaque export Excel contient une feuille
« Contrôle » (lignes, totaux par colonne numérique, empreinte SHA-256 de la source) pour prouver sa fidélité.
"""
import argparse
import csv
import datetime as dt
import hashlib
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "qualite-donnees" / "scripts"))
from qualite import load, num  # noqa: E402


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def numeric_cols(head, rows):
    out = []
    for i, _ in enumerate(head):
        vals = [r[i] for r in rows if i < len(r) and r[i] not in (None, "")]
        if vals and sum(num(v) is not None for v in vals) >= 0.9 * len(vals):
            out.append(i)
    return out


def totals(head, rows):
    return {head[i]: round(sum(num(r[i]) or 0 for r in rows if i < len(r)), 2) for i in numeric_cols(head, rows)}


# --- Excel mis en forme -------------------------------------------------------------------------
def to_excel(src: Path, dest: Path, sheet, titre, classe):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    head, rows = load(src, sheet)
    ncols = numeric_cols(head, rows)
    wb = Workbook()
    ws = wb.active
    ws.title = (titre or src.stem)[:31]
    start = 1
    if titre:
        ws.cell(1, 1, titre).font = Font(bold=True, size=14, color="1F3864")
        ws.cell(2, 1, f"{classe} — édité le {dt.date.today():%d/%m/%Y}").font = Font(italic=True, size=9, color="7A0010")
        start = 4
    thin = Side(style="thin", color="BFBFBF")
    for j, h in enumerate(head, 1):
        c = ws.cell(start, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)
    for i, r in enumerate(rows, start + 1):
        for j in range(len(head)):
            v = r[j] if j < len(r) else None
            if (j) in ncols and isinstance(v, str):
                v = num(v)
            c = ws.cell(i, j + 1, v)
            if isinstance(v, str) and v[:1] == "=":     # texte, jamais formule (anti-injection)
                c.data_type = "s"
            if j in ncols:
                c.number_format = "# ##0.00;[Red]-# ##0.00" if any(
                    isinstance(x[j], float) and x[j] % 1 for x in rows if j < len(x)) else "# ##0;[Red]-# ##0"
            elif isinstance(v, (dt.date, dt.datetime)):
                c.number_format = "dd/mm/yyyy"
            if (i - start) % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F5FA")
    last = start + len(rows)
    if ncols:                                   # ligne de total avec formules (recalculées par Excel)
        ws.cell(last + 1, 1, "TOTAL").font = Font(bold=True)
        for j in ncols:
            col = get_column_letter(j + 1)
            c = ws.cell(last + 1, j + 1, f"=SUBTOTAL(9,{col}{start + 1}:{col}{last})")
            c.font = Font(bold=True)
            c.number_format = "# ##0;[Red]-# ##0"
            c.border = Border(top=Side(style="double"))
    for j, h in enumerate(head, 1):             # largeur automatique bornée
        width = max([len(str(h))] + [len(str(r[j - 1])) for r in rows[:2000] if j - 1 < len(r) and r[j - 1] is not None])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 50)
    ws.freeze_panes = ws.cell(start + 1, 1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(head))}{last}"
    ws.print_title_rows = f"{start}:{start}"
    ws.page_setup.orientation = "landscape" if len(head) > 6 else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.center.text = classe
    ws.oddFooter.right.text = "Page &P / &N"
    ctl = wb.create_sheet("Contrôle")
    info = [("Source", src.name), ("SHA-256 source", sha256(src)), ("Lignes", len(rows)), ("Colonnes", len(head)),
            ("Généré le", dt.datetime.now().strftime("%d/%m/%Y %H:%M"))] + \
           [(f"Total {k}", v) for k, v in totals(head, rows).items()]
    for i, (k, v) in enumerate(info, 1):
        ctl.cell(i, 1, k).font = Font(bold=True)
        ctl.cell(i, 2, v)
    ctl.column_dimensions["A"].width, ctl.column_dimensions["B"].width = 28, 70
    wb.properties.creator = wb.properties.lastModifiedBy = ""
    wb.properties.title = titre or ""
    wb.save(dest)
    print(f"Excel : {dest}  ({len(rows)} lignes, {len(ncols)} colonnes chiffrées, feuille « Contrôle » ajoutée)")


# --- CSV compatible Excel français --------------------------------------------------------------
def to_csv(src: Path, dest: Path, sheet, sep, dec):
    head, rows = load(src, sheet)
    ncols = set(numeric_cols(head, rows))

    def fmt(j, v):
        if v is None:
            return ""
        if isinstance(v, (dt.date, dt.datetime)):
            return v.strftime("%d/%m/%Y")
        if j in ncols and num(v) is not None:
            n = num(v)
            s = f"{n:.2f}".rstrip("0").rstrip(".") if n % 1 else str(int(n))
            return s.replace(".", dec)
        s = str(v)
        return "'" + s if s[:1] in ("=", "+", "-", "@") and j not in ncols else s   # anti-injection de formule
    with open(dest, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=sep, quoting=csv.QUOTE_MINIMAL)
        w.writerow(head)
        for r in rows:
            w.writerow([fmt(j, r[j] if j < len(r) else None) for j in range(len(head))])
    print(f"CSV : {dest}  ({len(rows)} lignes, UTF-8 avec BOM, séparateur « {sep} », décimale « {dec} »)")


# --- rapprochement source / export --------------------------------------------------------------
def controle(src: Path, exp: Path, cle):
    h1, r1 = load(src, None)
    h2, r2 = load(exp, None)
    if h1[0] not in h2:                         # export avec titre : retrouver la ligne d'en-tête
        rows = [h2] + r2
        k = next((i for i, r in enumerate(rows) if r and str(r[0]).strip() == h1[0]), None)
        if k is None:
            sys.exit("En-tête de la source introuvable dans l'export.")
        h2, r2 = [str(x).strip() if x is not None else "" for x in rows[k]], rows[k + 1:]
    r2 = [r for r in r2 if str(r[0]).strip().upper() != "TOTAL"]
    ok = True
    print(f"Lignes : source {len(r1)} / export {len(r2)}", "✔" if len(r1) == len(r2) else "✘")
    ok &= len(r1) == len(r2)
    miss = [h for h in h1 if h not in h2]
    if miss:
        print("✘ colonnes absentes de l'export :", ", ".join(miss))
        ok = False
    t1, t2 = totals(h1, r1), totals(h2, r2)
    for k, v in t1.items():
        w = t2.get(k)
        good = w is not None and abs(v - w) <= max(0.01, abs(v) * 1e-9)
        ok &= good
        print(f"Total {k} : {v:,.2f} / {w if w is None else f'{w:,.2f}'}".replace(",", " "), "✔" if good else "✘")
    if cle and cle in h1 and cle in h2:
        k1 = {str(r[h1.index(cle)]).strip() for r in r1}
        k2 = {str(r[h2.index(cle)]).strip() for r in r2}
        print(f"Clés « {cle} » : {len(k1 - k2)} manquante(s), {len(k2 - k1)} en trop", "✔" if k1 == k2 else "✘")
        ok &= k1 == k2
    print("EXPORT FIDÈLE À LA SOURCE." if ok else "ÉCARTS DÉTECTÉS : ne pas diffuser avant explication.")
    sys.exit(0 if ok else 1)


# --- métadonnées Office -------------------------------------------------------------------------
META_RX = r"<(dc:creator|cp:lastModifiedBy|dc:title|dc:subject|dc:description|cp:keywords|Company|Manager|cp:category)>(.*?)</\1>"


def metadonnees(path: Path, clean: bool):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        xml = "".join(z.read(n).decode("utf-8", "ignore") for n in ("docProps/core.xml", "docProps/app.xml") if n in names)
        found = [(t, v) for t, v in re.findall(META_RX, xml, re.S) if v.strip()]
        extra = [n for n in names if re.search(r"comments|people\.xml|customXml|threadedComments", n, re.I)]
    for t, v in found:
        print(f"• {t.split(':')[-1]} : {v[:60]}")
    for n in extra:
        print(f"• élément sensible possible : {n}")
    if not found and not extra:
        print("Aucune métadonnée personnelle.")
        return
    if not clean:
        print("→ relancer avec --nettoyer pour produire une copie nettoyée (l'original est conservé).")
        return
    dest = path.with_name(path.stem + "_nettoye" + path.suffix)
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in ("docProps/core.xml", "docProps/app.xml"):
                data = re.sub(META_RX, lambda m: f"<{m.group(1)}></{m.group(1)}>", data.decode("utf-8"), flags=re.S).encode()
            zout.writestr(item, data)
    print(f"Copie nettoyée (auteur, société, titre…) : {dest}")
    if extra:
        print("⚠  commentaires / relecteurs présents : à supprimer dans Office (Révision → Supprimer tout), "
              "le script ne retire pas de contenu.")


def main():
    ap = argparse.ArgumentParser(description="Exports propres et contrôlés")
    ap.add_argument("action", choices=["excel", "csv", "controle", "metadonnees"])
    ap.add_argument("fichiers", nargs="+", type=Path)
    ap.add_argument("--sortie", type=Path)
    ap.add_argument("--feuille")
    ap.add_argument("--titre", default="")
    ap.add_argument("--classe", default="CONFIDENTIEL – USAGE INTERNE")
    ap.add_argument("--sep", default=";")
    ap.add_argument("--decimale", default=",")
    ap.add_argument("--cle")
    ap.add_argument("--nettoyer", action="store_true")
    a = ap.parse_args()
    src = a.fichiers[0]
    if not src.exists():
        sys.exit(f"Introuvable : {src}")
    if a.action in ("excel", "csv"):
        dest = a.sortie or src.with_name(src.stem + "_export" + (".xlsx" if a.action == "excel" else ".csv"))
        if dest.resolve() == src.resolve():
            sys.exit("La sortie doit être différente de la source.")
        if dest.exists():
            shutil.copy2(dest, Path(tempfile.gettempdir()) / (dest.name + ".precedent"))
        (to_excel(src, dest, a.feuille, a.titre, a.classe) if a.action == "excel"
         else to_csv(src, dest, a.feuille, a.sep, a.decimale))
    elif a.action == "controle":
        if len(a.fichiers) < 2:
            sys.exit("Donner la source puis l'export.")
        controle(src, a.fichiers[1], a.cle)
    else:
        metadonnees(src, a.nettoyer)


if __name__ == "__main__":
    main()
