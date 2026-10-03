#!/usr/bin/env python3
"""Garde-fou d'export : contrôle, pseudonymisation, chiffrement et journal de tout fichier qui sort.

    python export_guard.py scan  <fichier>                         # secrets + données personnelles
    python export_guard.py pseudo <fichier.xlsx|csv> --cols "Nom,Téléphone,Compte" [--sortie f]
    python export_guard.py chiffrer <fichier> [<fichier>…] --sortie envoi.zip   # ZIP AES-256, mot de passe saisi masqué
    python export_guard.py journal                                 # derniers exports

Formats lus : .xlsx .csv .txt .md .html .json .docx .pptx. Les valeurs trouvées sont toujours masquées.
La pseudonymisation remplace chaque valeur par un code stable (HMAC-SHA256 avec une clé locale), donc les
jointures et totaux restent possibles sans révéler les personnes.
"""
import csv
import datetime as dt
import getpass
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent.parent / "coffre-fort" / "scripts"))
from patterns import COMPILED, COMPILED_PERSONAL, luhn, mask, plausible  # noqa: E402

VAULT = Path(os.getenv("COFFRE_FORT_DIR", Path.home() / ".coffre-fort"))
JOURNAL = VAULT / "journal_exports.csv"


# --- lecture du texte de n'importe quel fichier courant ----------------------------------------
def text_of(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                out.append("\t".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    if ext in (".docx", ".pptx"):
        with zipfile.ZipFile(path) as z:
            xml = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
        return re.sub(r"<[^>]+>", " ", xml)
    return path.read_text("utf-8", errors="ignore")


def scan(path: Path) -> list:
    found = []
    for n, line in enumerate(text_of(path).splitlines(), 1):
        for name, rx in COMPILED + COMPILED_PERSONAL:
            for m in rx.finditer(line):
                if (name == "Carte bancaire" and not luhn(m.group(0))) or not plausible(name, m):
                    continue
                found.append((n, name, mask(m.group(0))))
    return found


# --- pseudonymisation ---------------------------------------------------------------------------
def _key() -> bytes:
    VAULT.mkdir(parents=True, exist_ok=True)
    k = VAULT / "cle_pseudonymes.bin"
    if not k.exists():
        k.write_bytes(secrets.token_bytes(32))
        try:
            os.chmod(k, 0o600)
        except OSError:
            pass
    return k.read_bytes()


def pseudo_value(v, key: bytes, prefix: str) -> str:
    if v is None or str(v).strip() == "":
        return v
    return f"{prefix}-{hmac.new(key, str(v).strip().upper().encode(), hashlib.sha256).hexdigest()[:10].upper()}"


def pseudonymize(src: Path, cols: list[str], dest: Path) -> Path:
    key = _key()
    wanted = {c.strip().lower() for c in cols}
    if src.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(src)
        for ws in wb.worksheets:
            head = {i: str(c.value).strip().lower() for i, c in enumerate(ws[1]) if c.value is not None}
            idx = [i for i, h in head.items() if h in wanted]
            for row in ws.iter_rows(min_row=2):
                for i in idx:
                    row[i].value = pseudo_value(row[i].value, key, head[i][:3].upper())
        wb.save(dest)
    else:
        raw = src.read_text("utf-8-sig")
        dialect = csv.Sniffer().sniff(raw.splitlines()[0], delimiters=",;\t")
        rows = list(csv.reader(io.StringIO(raw), dialect))
        head = [h.strip().lower() for h in rows[0]]
        idx = [i for i, h in enumerate(head) if h in wanted]
        for r in rows[1:]:
            for i in idx:
                if i < len(r):
                    r[i] = pseudo_value(r[i], key, head[i][:3].upper())
        with open(dest, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f, dialect).writerows(rows)
    found_cols = set(head.values()) if isinstance(head, dict) else set(head)
    missing = wanted - found_cols
    if missing:
        print("⚠  colonnes introuvables :", ", ".join(sorted(missing)))
    return dest


# --- chiffrement et journal ---------------------------------------------------------------------
def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def log(action: str, files: list[Path], note: str = ""):
    VAULT.mkdir(parents=True, exist_ok=True)
    new = not JOURNAL.exists()
    with open(JOURNAL, "a", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        if new:
            w.writerow(["date", "action", "fichier", "sha256", "note"])
        for p in files:
            w.writerow([dt.datetime.now().isoformat(timespec="seconds"), action, p.name, sha256(p), note])


def encrypt(files: list[Path], dest: Path) -> Path:
    import pyzipper
    pw = getpass.getpass("Mot de passe du ZIP (≥ 12 caractères, à transmettre par un AUTRE canal) : ")
    if len(pw) < 12 or pw != getpass.getpass("Confirmez : "):
        sys.exit("Mot de passe trop court ou différent : rien n'a été créé.")
    with pyzipper.AESZipFile(dest, "w", compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as z:
        z.setpassword(pw.encode())
        for p in files:
            z.write(p, p.name)
    del pw
    return dest


def main():
    import argparse
    ap = argparse.ArgumentParser(description="Garde-fou d'export")
    ap.add_argument("action", choices=["scan", "pseudo", "chiffrer", "journal"])
    ap.add_argument("fichiers", nargs="*", type=Path)
    ap.add_argument("--cols", default="", help="pseudo : colonnes à pseudonymiser, séparées par des virgules")
    ap.add_argument("--sortie", type=Path)
    ap.add_argument("--destinataire", default="", help="chiffrer : pour le journal")
    ap.add_argument("--force", action="store_true", help="chiffrer malgré des éléments sensibles (accord explicite)")
    a = ap.parse_args()
    if a.action == "scan":
        bad = 0
        for f in a.fichiers:
            res = scan(f)
            bad += len(res)
            for n, name, val in res[:200]:
                print(f"⚠  {f.name}:{n}  {name}  {val}")
            print(f"{f.name} : {len(res)} élément(s) sensible(s)" if res else f"{f.name} : rien de sensible détecté")
        sys.exit(1 if bad else 0)
    if a.action == "pseudo":
        src = a.fichiers[0]
        dest = a.sortie or src.with_name(src.stem + "_pseudonymise" + src.suffix)
        pseudonymize(src, a.cols.split(","), dest)
        log("pseudonymisation", [dest], f"colonnes={a.cols}")
        print("Fichier pseudonymisé :", dest)
    elif a.action == "chiffrer":
        dest = a.sortie or Path("export_chiffre.zip")
        for f in a.fichiers:                                  # contrôle obligatoire avant de chiffrer
            res = [r for r in scan(f) if r[1] not in ("E-mail", "Téléphone")]
            if res and not a.force:
                sys.exit(f"Refusé : {f.name} contient {len(res)} élément(s) sensible(s) (secrets/IBAN/cartes). "
                         "Pseudonymisez, ou relancez avec --force après accord explicite.")
        encrypt(a.fichiers, dest)
        log("export chiffré", [dest], a.destinataire)
        print(f"ZIP chiffré AES-256 : {dest}  (sha256 {sha256(dest)[:16]}…)")
    else:
        print(JOURNAL.read_text("utf-8") if JOURNAL.exists() else "Journal vide.")


if __name__ == "__main__":
    main()
