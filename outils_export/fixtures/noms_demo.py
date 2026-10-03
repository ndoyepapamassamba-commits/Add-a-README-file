"""Variante de démonstration des arrêtés synthétiques : dix grands clients reçoivent des noms de groupe fictifs.

M_01..M_09.xlsx : arrêtés entièrement synthétiques (SOCIETE 0000…, comptes L0000000…), aucune donnée réelle.
N_01..N_09.xlsx : mêmes arrêtés, avec trois vraies paires de groupe (BAOBAB, KORA, TERANGA) et deux pièges qui
ne partagent qu'un mot géographique (SAHEL, NIAYES), pour tester les « bénéficiaires potentiellement liés ».
Usage : python3 noms_demo.py   (dans ce dossier)
"""
import openpyxl

NOMS = {"100000378": "BAOBAB NEGOCE SA", "100000355": "BAOBAB NEGOCE DISTRIBUTION SARL",
        "100000388": "KORA TRANSPORTS ET FILS", "100000370": "KORA LOGISTIQUE ET TRANSIT",
        "100000353": "SAHEL CIMENTS INDUSTRIE SA", "100000461": "SAHEL PECHE ARMEMENT",
        "100000483": "TERANGA BATIMENT TRAVAUX PUBLICS", "100000454": "TERANGA IMMOBILIER PROMOTION",
        "100000561": "NIAYES HORTICULTURE EXPORT", "100000364": "COOPERATIVE AGRICOLE DES NIAYES"}

for k in range(1, 10):
    wb = openpyxl.load_workbook(f"M_0{k}.xlsx")
    ws = wb["ACTE 7"]
    entete = [c.value for c in ws[1]]
    cc, rl = entete.index("Code Client"), entete.index("Relationship")
    for ligne in ws.iter_rows(min_row=2):
        if str(ligne[cc].value) in NOMS:
            ligne[rl].value = NOMS[str(ligne[cc].value)]
    wb.save(f"N_0{k}.xlsx")
    print(f"N_0{k}.xlsx")
